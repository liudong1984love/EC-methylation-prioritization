# 49_v99_revisions.py — Manuscript v9.8 -> v9.9
# Addresses reviewer round 5 (14 points): R version detail, pipeline wording,
# ZSCAN12 substitute borderline probe, P95 definition, cross-hybridization wording,
# literature-search hedging, comb-p role, sample-size caveat, reviewer-tone removal,
# abstract aim softening, pan-platform fix, Figure 3 three-tier replacement, compression.
import copy
from docx import Document
from docx.shared import Inches
from pathlib import Path

BASE = Path("C:/Users/ld/WorkBuddy/2026-08-19-15-50-19")
SRC = BASE / "Manuscript_EC_methylation_panel_v9.8.docx"
DST = BASE / "Manuscript_EC_methylation_panel_v9.9.docx"

doc = Document(SRC)

def set_text(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)

log = []

def sub(old, new, tag, count_expected=1):
    n = 0
    for p in doc.paragraphs:
        if old in p.text:
            set_text(p, p.text.replace(old, new))
            n += 1
    log.append(("OK  " if n == count_expected else f"WARN n={n}") + f" [{tag}] {old[:70]}")

def para_rewrite(anchor, new_text, tag):
    for p in doc.paragraphs:
        if p.text.strip().startswith(anchor):
            set_text(p, new_text)
            log.append(f"OK   [{tag}] paragraph rewritten")
            return
    log.append(f"MISS [{tag}] {anchor[:60]}")

def wc():
    return sum(len(p.text.split()) for p in doc.paragraphs)

w0 = wc()

# ---------------- P12: abstract aim softening ----------------
sub("We aimed to build and rigorously triage methylation candidates specifically for premenopausal EC.",
    "We aimed to build and rigorously triage methylation candidates for subsequent validation in premenopausal EC populations.",
    "P12-abstract-aim")

# ---------------- P7: abstract pipeline wording ----------------
sub("with a unified minfi/limma pipeline and comb-p-like region merging",
    "with harmonised downstream differential-methylation criteria and comb-p-like region merging",
    "P7-abstract-pipeline")

# ---------------- P2: abstract background comparison wording ----------------
sub("Background estimates were higher in the benign endometrial cohort than in healthy donors (approximately 1.2- to 1.6-fold on a uniform pipeline).",
    "Background estimates were higher in the intended-use-proxy cohort than in a small cancer-free endometrial reference cohort (approximately 1.2- to 1.6-fold on the same platform and reported normalisation framework).",
    "P2-abstract-bg")

# ---------------- P3: abstract substitute probes ----------------
sub("three EPIC-compatible probes within the same DMR passed all background gates at the region level",
    "three EPIC-compatible probes within the same DMR passed (two) or marginally exceeded (one, borderline) the background gates at the region level",
    "P3-abstract-substitutes")

# ---------------- P7: methods data sources (TCGA processing) ----------------
sub("TCGA-UCEC methylation data and clinical variables including menopause status were obtained from UCSC Xena.",
    "TCGA-UCEC methylation data were obtained from UCSC Xena as precomputed \u03b2 values (raw IDATs were not re-processed for TCGA, so upstream normalisation differs from the minfi/noob processing of GSE67116 and GSE73949); clinical variables including menopause status were obtained from the same source.",
    "P7-methods-tcga")

# ---------------- P8 + P5: DMR methods ----------------
sub("DMRs were defined with a deterministic comb-p-like algorithm [18]: same-direction significant probes (FDR<0.05) within 500 bp were merged, regions with \u22653 probes retained, and region-level Stouffer P values computed.",
    "DMRs were defined with a deterministic comb-p-like algorithm [18]: same-direction significant probes (FDR<0.05) within 500 bp were merged, regions with \u22653 probes retained, and region-level Stouffer P values computed. Region-level P values were used only for candidate generation; final selection was driven by effect size, tumour positivity and background gates, and these P values are not interpreted as calibrated confirmatory evidence.",
    "P8-combp-role")

sub("Both BOLL-region probes map uniquely to the BOLL locus (chr2:197,786,574 and chr2:197,786,304; mapQ 60) in the Zhou et al. hg38 re-annotation [28], excluding cross-hybridization as an explanation for their discordant behaviour.",
    "Both BOLL-region probes map uniquely to the BOLL locus (chr2:197,786,574 and chr2:197,786,304; mapQ 60) and carry no common SNPs at the CpG or single-base-extension sites in the Zhou et al. hg38 re-annotation [28], making cross-mapping an unlikely explanation for their discordant behaviour; residual probe-sequence cross-reactivity cannot be excluded by annotation alone.",
    "P5-crosshyb")

# ---------------- P1 + P4 + P7: statistical analysis ----------------
sub("Analyses used R 4.6.0 and Python 3.13.",
    "Analyses used R 4.6.0 (2026-04-24, x86_64-w64-mingw32; minfi 1.58.0, limma 3.68.4, bumphunter 1.54.0) and Python 3.13. Background P95 was computed with linear-interpolation quantiles (equivalent to R quantile type 7).",
    "P1-P4-stats")

sub("Discovery-layer \u03b2 values were noob-normalised, whereas background triage used GMQN-normalised \u03b2 values",
    "Discovery-layer \u03b2 values combined Xena-precomputed TCGA data with noob-normalised GSE67116 data, whereas background triage used GMQN-normalised \u03b2 values",
    "P7-stats-discovery")

# ---------------- P2: results cohort-dependence wording ----------------
sub("on a uniform EPIC/GMQN pipeline the benign cohort remained consistently elevated",
    "in an EPIC/GMQN-aligned comparison (same platform and reported normalisation framework, different cohorts and centres) the benign cohort remained consistently elevated",
    "P2-results-bg")

# ---------------- P6: literature convergence ----------------
sub("and none of our novel candidates, including BOLL, appeared in that study or elsewhere in the EC detection literature.",
    "and none of our novel candidates appeared in that study; no EC-detection report for BOLL was identified in our literature search.",
    "P6-litconvergence")

# ---------------- P11: remove reviewer-tone in threshold paragraph ----------------
sub("We therefore correct the earlier blanket statement that both candidates tolerate a 0.10\u20130.20 gate: rank stability and all-gate compliance are distinct properties, gate choice materially affects BOLL's pass/fail status, and final cut-offs must be set empirically in the validation cohort.",
    "Rank stability across gates is therefore distinct from strict all-gate compliance; gate choice materially affects BOLL's pass/fail status, and final cut-offs must be set empirically in the validation cohort.",
    "P11-reviewer-tone")

# ---------------- P3 + P13: ZSCAN12 substitutes paragraph ----------------
para_rewrite("To address the EPIC-platform absence of cg27577527",
    "To address the EPIC-platform absence of cg27577527, we screened all nine EPIC-manifest probes spanning the same ZSCAN12 DMR (hg38 chr6:28,399,501\u201328,400,362, containing cg27577527 at 28,399,766) against the decisive intended-use-proxy background (GSE223817) and the external tumour cohorts. Three probes\u2014cg23164203, cg20275132 and cg25666433\u2014passed the P95 \u03b2\u22640.15 gate in the approximating cohort (P95 0.088\u20130.120; positivity \u22640.3%) with low cervical (0.060\u20130.095), blood (0.080\u20130.092) and cycle-drift (0.023\u20130.029) backgrounds while retaining high tumour positivity (TCGA 86.8\u201390.3% overall, 89.3\u201396.4% premenopausal; GSE136791 carcinoma 91.3\u201395.7%; GSE155760 65.2\u201373.9%; per-cohort metrics in Supplementary Data). At the healthy-endometrium dimension (n=17), cg23164203 (P95 0.143, bootstrap 95% CI 0.123\u20130.144) and cg20275132 (0.149, 95% CI 0.129\u20130.154) passed by point estimate, whereas cg25666433 was borderline (0.151, 95% CI 0.124\u20130.155, crossing the gate). The ZSCAN12 region therefore remains a candidate region with both 450K- and EPIC-compatible readouts, its approximating-cohort compliance being supported at the region level by EPIC-compatible probes that can be incorporated into the targeted assay.",
    "P3-P13-substitutes")

# ---------------- P6: discussion BOLL novelty ----------------
sub("BOLL (cg24589459) is, to our knowledge, unreported for EC detection.",
    "No EC-detection report for BOLL (cg24589459) was identified in our literature search.",
    "P6-boll-novelty")

sub("We caution, however, that absence from the literature may also reflect unpublished negative results; the first-mover claim must be settled by wet-lab data, ideally in the same cohorts that will adjudicate the literature anchors.",
    "Absence from the literature may also reflect unpublished negative results; whether this novelty carries diagnostic value must be settled by wet-lab data, ideally in the same cohorts that will adjudicate the literature anchors.",
    "P6-firstmover")

# ---------------- P2: discussion cohort-dependence ----------------
sub("1.2- to 1.6-fold on a uniform pipeline, up to two- to three-fold raw",
    "1.2- to 1.6-fold in the EPIC/GMQN-aligned comparison, up to two- to three-fold raw",
    "P2-discussion-bg")

# ---------------- P4: limitations bootstrap caveat ----------------
sub("The healthy-donor background estimates rest on 17 samples, so P95 statistics are unstable (bootstrap 95% CI for cg07495363 spans 0.099\u20130.203);",
    "The healthy-donor background estimates rest on 17 samples, so P95 statistics are unstable (bootstrap 95% CI for cg07495363 spans 0.099\u20130.203), and bootstrap intervals for extreme quantiles in small cohorts are driven by one or two extreme observations;",
    "P4-bootstrap-caveat")

# ---------------- P10: sample-size caveat ----------------
sub("a definitive cohort of approximately 140 carcinomas and 200 benign controls would narrow both to \u00b15%.",
    "a definitive cohort of approximately 140 carcinomas and 200 benign controls would narrow both to \u00b15%. These are approximate single-proportion precision targets, not a formal diagnostic-accuracy design; they should be refined with allowance for assay failure, separate threshold-development and validation sets, and prespecified handling of atypical hyperplasia.",
    "P10-samplesize")

# ---------------- P14: Figure 3 legend ----------------
sub("Green: P95 \u03b2 \u22640.15 in both cervix and blood (favourable or borderline-favourable point-estimate in silico background profile); red: compartment-restricted.",
    "Green: P95 \u03b2 \u22640.15 in both cervix and blood (clear pass); yellow: point-estimate pass with bootstrap confidence interval crossing the gate (borderline; BOLL anchor); red: compartment-restricted (P95 \u03b2 >0.15).",
    "P14-fig3-legend")

# ---------------- P14: replace Figure 3 image ----------------
replaced_img = False
for i, p in enumerate(doc.paragraphs):
    blips = p._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
    if blips and i + 1 < len(doc.paragraphs) and doc.paragraphs[i + 1].text.startswith("Figure 3."):
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        p.add_run().add_picture(str(BASE / "results" / "v99_figure3_compartment.png"), width=Inches(5.8))
        replaced_img = True
        break
log.append(("OK  " if replaced_img else "MISS") + " [P14-fig3-image] replaced")

# ---------------- P11: compression pass ----------------
para_rewrite("Three candidates showed favourable (ZSCAN12, ARL5C)",
    "Three candidates showed favourable (ZSCAN12, ARL5C) or borderline-favourable (BOLL) point-estimate background profiles across all assessed compartments (Table 7); these profiles reflect purified-reference backgrounds and do not establish suitability for real self-collected samples, in which tumour-DNA fraction, cellular mixing and degradation dominate. CDO1 showed a dual-compartment profile (cervical 0.055, blood 0.080) consistent with its use in both cervical and intrauterine formats. The approved-kit genes AJAP1 and GALR1 did not meet our prespecified probe-level background criteria for cervical or self-sampling applications because of blood-cell background, consistent with their deployment via an intrauterine brush [10]; this probe-level audit does not evaluate the exact CpGs, amplicons or classification algorithms used in commercial or published assays.",
    "compress-compartment")

para_rewrite("Despite this progress, three methodological gaps remain",
    "Despite this progress, three methodological gaps remain. First, discovery and validation cohorts are dominated by postmenopausal women; CISENDO is, to our knowledge, the first menopause-stratified validation, and no marker set has been designed for premenopausal physiology. Second, benign controls are usually adjacent tissue or healthy donors, whereas intended-use populations are symptomatic women with abnormal uterine bleeding (AUB); the extent to which marker background differs between these control types is unknown. Third, markers are rarely evaluated for the sampling compartment in which a test will run: cervical scraping or self-sampling dilutes tumour DNA and adds cervical and blood background, and urine or ctDNA formats are constrained primarily by blood-cell background [12]. Pan-cancer loci such as ZNF154 illustrate the risk of assigning a strong tumour signal to a compartment where its background is unacceptable [11], and normal endometrial methylation is itself cycle-dependent [13,14].",
    "compress-intro-gaps")

para_rewrite("The cohort dependence we quantified",
    "The cohort dependence we quantified (background inflation in the intended-use-proxy cohort versus healthy endometrium; 1.2- to 1.6-fold in the EPIC/GMQN-aligned comparison, up to two- to three-fold raw) offers a concrete explanation for the attrition of methylation markers during translation and argues that benign controls must be drawn from the intended-use population. Public data remain devoid of premenopausal benign AUB methylation profiles (polyps, leiomyomas, adenomyosis), a gap only purpose-built cohorts can fill. More broadly, strong tumour-associated methylation does not guarantee suitability for a particular target population or sampling compartment.",
    "compress-discussion-cohort")

para_rewrite("The audit of the internally derived set reinforces",
    "The audit of the internally derived set reinforces the same lesson from the opposite direction: four of five markers that appeared excellent by tumour signal alone (including the pan-cancer locus ZNF154 [11]) failed in population- or compartment-relevant backgrounds. Markers are not wrong in absolute terms, only misassigned to a population or compartment; the audit shares the TCGA discovery data and is therefore not independent validation.",
    "compress-audit-discussion")

para_rewrite("We next tested whether panel loci show a quantitative methylation gradient",
    "We next tested whether panel loci show a quantitative methylation gradient across the neoplastic continuum. Within GSE136791 (EPIC), methylation was consistently higher in carcinoma (n=69) than in hyperplasia (n=34) for all 21 evaluable panel probes (median mean \u03b2 0.665 versus 0.601); this within-cohort contrast is the primary evidence. Extending the gradient to benign endometrium is descriptive, because the benign reference (13 endometrial controls from GSE155760, same platform) came from a different cohort; with that reference, mean \u03b2 increased monotonically from benign through hyperplasia to carcinoma for all 21 probes (median 0.152, 0.601, 0.665; both BOLL probes followed the same ordering). The carcinoma>hyperplasia ordering was reproduced in GSE67116 (450K; all 23 probes; median 0.533 versus 0.327), which contributed to discovery and is therefore not independent evidence. Hyperplasia histology (atypical versus non-atypical) is not annotated in either cohort. Because the intended use is triage of abnormal uterine bleeding\u2014in which detection of atypical hyperplasia is clinically valuable rather than a false positive\u2014we interpret hyperplasia positivity as detection of the neoplastic continuum and explicitly avoid carcinoma-specificity claims.",
    "compress-gradient")

doc.save(DST)
w1 = wc()
log.append(f"word count: {w0} -> {w1} ({100*(w1-w0)/w0:+.1f}%)")
print("\n".join(log))
