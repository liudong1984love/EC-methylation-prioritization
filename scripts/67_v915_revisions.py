# -*- coding: utf-8 -*-
"""v9.15: 9th review round — (1) drop direct atypical-hyperplasia performance claim;
(2) GSE223817 no longer called 'intended-use premenopausal population' in abstract;
(3) unified adjusted-model description (design formula + beta-scale marginal means);
(4) author info NOT touched per user instruction.
Bonus: 'descriptively higher'; GitHub 'available' vs Zenodo 'permanently archived'."""
from docx import Document

SRC = r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.14.docx"
DST = r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.15.docx"

doc = Document(SRC)
log = []

def set_text(p, new):
    for r in list(p.runs):
        r.text = ""
    if p.runs:
        p.runs[0].text = new
    else:
        p.add_run(new)

def replace_in_para(anchor, old, new, must=True):
    n = 0
    for p in doc.paragraphs:
        if anchor in p.text and old in p.text:
            set_text(p, p.text.replace(old, new))
            n += 1
    log.append(("OK " if n >= 1 else "MISS ") + old[:60])
    if must and n == 0:
        raise SystemExit("MISS: " + old[:80])

def replace_para_by_prefix(prefix, new_text):
    n = 0
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            set_text(p, new_text)
            n += 1
    log.append(("OK " if n == 1 else "MISSx%d " % n) + prefix[:50])
    if n != 1:
        raise SystemExit("PREFIX MISS: " + prefix[:60])

def append_to_para(anchor, addition):
    n = 0
    for p in doc.paragraphs:
        if anchor in p.text and addition[:40] not in p.text:
            set_text(p, p.text.rstrip() + " " + addition)
            n += 1
    log.append(("OK " if n >= 1 else "MISS ") + ("APPEND " + anchor[:50]))
    if n == 0:
        raise SystemExit("APPEND MISS: " + anchor[:60])

# ---------- 1+2. Abstract: Methods (proxy wording) / Results (descriptively) / Conclusions (no AH claim) ----------
replace_para_by_prefix("Methods: We re-analysed",
"Methods: We re-analysed 450K data from TCGA-UCEC and GSE67116 with harmonised differential-methylation criteria and region merging. Candidates were triaged against five background dimensions spanning healthy endometrium, benign endometrium from a surgical cohort used as an intended-use proxy (individual menopausal status and AUB presentation unavailable), menstrual-cycle phases, cervical scrapes and purified blood cells. Panels were assembled by redundancy pruning and greedy coverage; the NMPA-kit genes and WID-qEC components were audited through the same funnel. Sensitivity analyses comprised menopause-stratified, patient-matched, subtype-restricted and adjusted models, cross-preprocessing checks and bootstrap intervals.")

replace_para_by_prefix("Results: Background estimates",
"Results: Background estimates were descriptively higher in the intended-use-proxy cohort than in a cancer-free endometrial reference cohort. Among the lead candidates, cg24589459 (BOLL) was the only probe evaluable across all five background dimensions that met the prespecified gate by point estimate, although its bootstrap confidence interval crossed the threshold (borderline). cg27577527 (ZSCAN12) passed the four probe-level dimensions with the most consistently low background; the intended-use-proxy dimension is not directly assessable (the probe is absent from the EPIC platform), and support there is indirect and region-level only. Corroboration in independent tumour-tissue cohorts—not the intended sampling compartments—was partial, yielding a post hoc two-probe region-level hypothesis for BOLL. Menopause-associated differences were non-significant but underpowered. Background estimates were pipeline-dependent, so the fixed β thresholds are heuristic. Estimates are partly discovery-derived upper bounds; the candidates are not a validated panel.")

replace_para_by_prefix("Conclusions: Marker background",
"Conclusions: Marker background is cohort- and compartment-dependent; benign controls should come from the target clinical population. BOLL and ZSCAN12 are prioritised for prospective evaluation in premenopausal abnormal-uterine-bleeding triage, where the intended target condition comprises endometrial cancer and atypical hyperplasia; performance for atypical hyperplasia cannot be established from the available datasets. Population matching was approximated, not achieved; confirmation requires wet-lab validation in a purpose-built premenopausal AUB cohort.")

replace_para_by_prefix("Background: Methylation-based",
"Background: Methylation-based endometrial cancer (EC) detection is maturing rapidly, with an NMPA-approved three-gene adjunctive kit, international programmes (WID-qEC; CISENDO) and a 2026 Chinese expert consensus. However, existing markers were discovered and validated predominantly in postmenopausal women, and marker background is rarely assessed in the symptomatic populations or sampling compartments in which tests are deployed. We aimed to build and triage methylation candidates for premenopausal EC validation.")

# ---------- 1b. Conclusions paragraph: same reframing ----------
replace_in_para("In conclusion, multi-cohort discovery",
    "prioritised the BOLL and ZSCAN12 regions for further evaluation in the detection of endometrial cancer and atypical hyperplasia during premenopausal AUB triage, and reconciled",
    "prioritised the BOLL and ZSCAN12 regions for prospective evaluation in premenopausal AUB triage—where the intended target condition comprises endometrial cancer and atypical hyperplasia, although performance for atypical hyperplasia cannot be established from the available datasets—and reconciled")

# ---------- 3. Adjusted-model description: single consistent design formula ----------
replace_in_para("Confounder sensitivity used limma models",
    "Confounder sensitivity used limma models on M-values adjusting for age and histological type (three-level tissue factor), an endometrioid-only age-adjusted model, and a tumour-only subtype-association model.",
    "Confounder sensitivity used limma models on M-values with the design ~ age + tissue_histology_group, where tissue_histology_group is a three-level factor (adjacent normal / endometrioid tumour / non-endometrioid tumour), so that normal tissue carries no separate histology term; an endometrioid-only age-adjusted model and a tumour-only subtype-association model were also fitted.")

replace_in_para("Covariate-adjusted sensitivity models used limma",
    "group plus age and histological type (endometrioid versus other; n=462)",
    "age plus tissue_histology_group (n=462)")

append_to_para("tumour-only molecular-subtype model (n=398",
    "Adjusted Δβ values were derived on the β scale as adjusted marginal-mean differences from a linear model with the same design (averaging the endometrioid and non-endometrioid tumour contrasts versus adjacent normal), not from M-value coefficients.")

# Results paragraph: annotate the 0.03 figure (phrase 'maximum change in Δβ 0.03' preserved)
replace_in_para("maximum change in Δβ 0.03",
    "maximum change in Δβ 0.03)",
    "maximum change in Δβ 0.03, β-scale adjusted marginal means)")

# ---------- bonus 1. 'descriptively higher' for cross-cohort contrasts ----------
replace_in_para("were higher than in healthy premenopausal donors",
    "were higher than in healthy premenopausal donors",
    "were descriptively higher than in healthy premenopausal donors")

# ---------- bonus 2. GitHub available / Zenodo permanently archived ----------
replace_in_para("Data and code availability",
    "are permanently archived on GitHub",
    "are available on GitHub")
replace_in_para("Data and code availability",
    ") and Zenodo (version v1.0.1",
    ") and permanently archived in Zenodo (version v1.0.1")

doc.save(DST)
print("\n".join(log))
print("saved", DST)

# abstract + total word counts
d2 = Document(DST)
abstract = [p.text for p in d2.paragraphs if p.text.strip().startswith(("Background:", "Methods:", "Results:", "Conclusions:"))][:4]
total = sum(len(t.split()) for t in abstract)
print("ABSTRACT WORDS:", total, [len(t.split()) for t in abstract])
allw = len("\n".join(p.text for p in d2.paragraphs if p.text.strip()).split())
print("TOTAL WORDS:", allw)
