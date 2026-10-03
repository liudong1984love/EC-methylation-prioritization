# -*- coding: utf-8 -*-
"""Manuscript v9.26 -> v9.27: three incremental insertions (Results x2, Discussion x1)."""
import zipfile, os

SRC = r'C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.26.docx'
DST = r'C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.27.docx'

B = '\u03b2'; LE = '\u2264'

EDITS = [
    # 1) Results: ZSCAN12 new 450K benign evidence (after the four-gate compliance disclaimer)
    ('We therefore do not claim four-gate compliance for ZSCAN12.',
     'We therefore do not claim four-gate compliance for ZSCAN12. '
     'A systematic extension of the background search (Supplementary Note S2.11) nevertheless identified two '
     'independent 450K cohorts in which this probe is assessable, providing its first direct benign-tissue '
     'assessments: no exceedances in five benign endometrial samples (P95 ' + B + ' 0.056) and one exceedance '
     'in 26 eutopic endometrium samples from endometriosis patients (P95 0.108)\u2014point-estimate-level '
     'support only, on which we still do not base a four-gate compliance claim.'),

    # 2) Results: CDO1 / BOLL partner low background in new cohorts (after the tier-B audit sentence)
    ('not as a gate-passing candidate (full audit, including per-probe values: Supplementary Note S1, Table S1).',
     'not as a gate-passing candidate (full audit, including per-probe values: Supplementary Note S1, Table S1). '
     'Consistent with a cohort-specific rather than universal elevation, both cg23180938 (CDO1) and the BOLL '
     'partner probe cg07495363 showed low background in two independently evaluated 450K cohorts (P95 ' + LE +
     ' 0.149 in benign endometrium, n=5, and in eutopic endometrium from endometriosis patients, n=26), in '
     'marked contrast to their high surgical-cohort background (P95 0.372 and 0.427; Supplementary Note S2.11).'),

    # 3) Discussion: systematic null result for public benign-lesion data
    ('a gap only purpose-built cohorts can fill.',
     'a gap only purpose-built cohorts can fill. A systematic search of the EWAS Data Hub (18 endometrium '
     'projects, 1,746 samples) confirmed this absence: no public 450K/EPIC cohorts exist for endometrial '
     'polyps, leiomyomas, adenomyosis or PCOS (Supplementary Note S2.11).'),
]

zin = zipfile.ZipFile(SRC)
xml = zin.read('word/document.xml').decode('utf-8')
for old, new in EDITS:
    n = xml.count(old)
    assert n == 1, f'anchor not unique ({n}): {old[:60]}'
    xml = xml.replace(old, new)
    print('applied:', old[:60], '...')

with zipfile.ZipFile(DST, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == 'word/document.xml':
            data = xml.encode('utf-8')
        zout.writestr(item, data)
zin.close()

z = zipfile.ZipFile(DST)
assert z.testzip() is None
import xml.etree.ElementTree as ET
ET.fromstring(z.read('word/document.xml'))
x = z.read('word/document.xml').decode('utf-8')
for probe in ['first direct benign-tissue', 'cohort-specific rather than universal', 'confirmed this absence']:
    assert probe in x, probe
import docx
d = docx.Document(DST)
print('paragraphs:', len(d.paragraphs), '| tables:', len(d.tables))
print('saved:', DST)
