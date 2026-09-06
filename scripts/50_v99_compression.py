# 50_v99_compression.py — v9.9 compression pass + ref27 access date
# Round-5 point 11: remove residual repetition, target ~5-6% cut on top of v9.8->v9.9
from docx import Document
from pathlib import Path

BASE = Path("C:/Users/ld/WorkBuddy/2026-08-19-15-50-19")
F = BASE / "Manuscript_EC_methylation_panel_v9.9.docx"
doc = Document(F)

def set_text(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)

log = []
def sub(old, new, tag):
    n = 0
    for p in doc.paragraphs:
        if old in p.text:
            set_text(p, p.text.replace(old, new)); n += 1
    log.append(("OK  " if n else "MISS") + f" [{tag}] x{n}")

def wc():
    return sum(len(p.text.split()) for p in doc.paragraphs)

w0 = wc()

# [17] methods contrasts
sub("Contrasts: TCGA tumour versus adjacent normal (unpaired, 431 versus 46, unadjusted for purity or subtype, because effect-size filtering rather than significance was the primary criterion) and GSE67116 primary tumour versus hyperplasia; cross-cohort consistency was assessed at the effect-size level (region mean \u0394\u03b2, same direction), because the eight-sample hyperplasia group limited formal DMR power.",
    "Contrasts were TCGA tumour versus adjacent normal (unpaired, 431 versus 46; unadjusted, effect-size filtering being the primary criterion) and GSE67116 primary tumour versus hyperplasia; cross-cohort consistency was assessed at the effect-size level (region mean \u0394\u03b2, same direction), the eight-sample hyperplasia group limiting formal DMR power.",
    "methods-contrasts")

# [11] intro gaps
sub("the extent to which marker background differs between these control types is unknown.",
    "whether marker background differs between these control types is unknown.",
    "intro-gaps")

# [34] cohort dependence
sub("We treat this as a descriptive cross-cohort comparison, because cohort, centre and residual processing effects cannot be fully separated and the reference comprises only 13 samples; the qualitative conclusion is platform-independent, whereas the raw two- to three-fold contrast was inflated by processing.",
    "We treat this as a descriptive cross-cohort comparison (cohort, centre and residual processing effects cannot be fully separated; reference n=13): the qualitative conclusion is platform-independent, the raw two- to three-fold contrast being inflated by processing.",
    "cohort-dependence")

# [37] final candidates
sub("absent from EPIC and therefore not assessable in the approximating cohort",
    "absent from EPIC, so not assessable in the approximating cohort",
    "candidates-epic")
sub("The high-sensitivity BOLL-region probe cg07495363 alone detected all 28 premenopausal tumours in silico (the anchor detected 25 of 28), with greedy coverage reaching 99.8% after two additions",
    "The BOLL partner probe cg07495363 alone detected all 28 premenopausal tumours in silico (anchor 25/28), greedy coverage reaching 99.8% after two additions",
    "candidates-greedy")
sub("of which 15 additionally passed the narrower cervix\u2013blood\u2013cycle triage gate, a distinct criterion from the union rule; full list, coordinates and primer annotations in Supplementary Data",
    "of which 15 additionally passed the narrower cervix\u2013blood\u2013cycle gate, a distinct criterion from the union rule; list, coordinates and primer annotations in Supplementary Data",
    "candidates-pool")

# [44] threshold paragraph: remove restated sentence
sub("Rank stability across gates is therefore distinct from strict all-gate compliance; gate choice materially affects BOLL's pass/fail status, and final cut-offs must be set empirically in the validation cohort.",
    "Gate choice therefore materially affects BOLL's pass/fail status; final cut-offs must be set empirically in the validation cohort.",
    "threshold-dedup")

# [45] menopause/subtype
sub("(mean paired \u0394\u03b2 0.46\u20130.70; \u226591% of pairs concordant; Wilcoxon FDR<0.01), showing that the signal persists in a patient-matched analysis and is unlikely to be explained solely by between-patient differences; residual confounding by tumour purity, cell composition, batch or field effects cannot be excluded.",
    "(mean paired \u0394\u03b2 0.46\u20130.70; \u226591% of pairs concordant; Wilcoxon FDR<0.01): the signal persists patient-matched and is unlikely to be explained solely by between-patient differences, although tumour purity, cell composition, batch and field effects remain as residual confounders.",
    "paired-compress")

# [47] external validation
sub("This dissociation\u2014reproducible tumour positivity of the anchor with an exceptionally clean control background\u2014indicates that cg24589459 behaves as a low-false-positive probe and cg07495363 as a high-sensitivity probe of the same region; based on these exploratory data, we propose prospectively specifying and locking a two-probe region-level rule before wet-lab validation.",
    "This dissociation indicates complementary behaviour within the region\u2014cg24589459 low-false-positive, cg07495363 high-sensitivity\u2014and, being exploratory, motivates our proposal to prospectively specify and lock a two-probe region-level rule before wet-lab validation.",
    "extval-dissociation")
sub("In GSE178610, positivity was strongly preservation-dependent (\u226580% for 14/21 probes in FFPE versus 2/21 in fresh-frozen; Supplementary Data), a pattern consistent with, but not specific for, fixation- or batch-related inflation;",
    "In GSE178610, positivity was preservation-dependent (\u226580% for 14/21 probes in FFPE versus 2/21 fresh-frozen; Supplementary Data), consistent with\u2014but not specific for\u2014fixation- or batch-related inflation;",
    "extval-ffpe")
sub("As a partial substitute for the EPIC-absent cg27577527, we assessed it in a small 450K cohort not used in discovery (GSE93589, n=9 EC): positive in 8/9 tumours (mean \u03b2=0.62).",
    "cg27577527 was additionally assessed in a small 450K cohort not used in discovery (GSE93589, n=9 EC): positive in 8/9 tumours (mean \u03b2=0.62).",
    "extval-93589")

# [51] BOLL expression paragraph -> single sentence
def para_rewrite(anchor, new_text, tag):
    for p in doc.paragraphs:
        if p.text.strip().startswith(anchor):
            set_text(p, new_text)
            log.append(f"OK   [{tag}] paragraph rewritten")
            return
    log.append(f"MISS [{tag}]")
para_rewrite("Methylation at cg24589459 correlated weakly",
    "BOLL expression is essentially absent from endometrial tissue irrespective of methylation state (172 TCGA tumours with paired RNA-seq; full analysis in Supplementary Data), so BOLL hypermethylation in EC is best interpreted as ectopic hypermethylation of a repressed germline-specific locus, not silencing of an expressed gene.",
    "boll-expression")

# [54] discussion convergence
sub("The convergence between our de novo ranking and independent wet-lab programmes is the strongest internal evidence of validity. Without using any prior gene list, our funnel re-identified the core of the most advanced international test (ZSCAN12 in WID-qEC [5,6]), the anchor of the only approved national kit (CDO1 [8,10]), and a component of the Dutch panel (GHSR [24]).",
    "Convergence with independent wet-lab programmes is the strongest internal evidence of validity: without any prior gene list, the funnel re-identified the core of WID-qEC (ZSCAN12 [5,6]), the anchor of the only approved national kit (CDO1 [8,10]) and a Dutch-panel component (GHSR [24]).",
    "disc-convergence")

# [55] BOLL discussion tail
sub("Two observations clarify the BOLL signal. First, both BOLL-region probes map uniquely to the locus [28], so the dissociation between the anchor (low background, moderate positivity) and its partner 270 bp away (higher background, high sensitivity) most likely reflects genuine regional methylation heterogeneity rather than probe artefact, reinforcing the need for region-level assay design. Second, the expression analysis indicates that BOLL is a passenger rather than a driver marker\u2014we make no claim of pathway involvement\u2014which bounds biological interpretation without reducing diagnostic utility.",
    "Both BOLL-region probes map uniquely to the locus [28], so the anchor/partner dissociation 270 bp apart most likely reflects genuine regional methylation heterogeneity rather than probe artefact, reinforcing the need for region-level assay design; and the expression analysis marks BOLL as a passenger rather than driver marker, bounding biological interpretation without reducing diagnostic utility.",
    "disc-boll")

# [57] audit discussion
sub("The audit of the internally derived set reinforces the same lesson from the opposite direction: four of five markers that appeared excellent by tumour signal alone (including the pan-cancer locus ZNF154 [11]) failed in population- or compartment-relevant backgrounds. Markers are not wrong in absolute terms, only misassigned to a population or compartment; the audit shares the TCGA discovery data and is therefore not independent validation.",
    "The internal-set audit reinforces the same lesson from the opposite direction: four of five markers excellent by tumour signal alone (including pan-cancer ZNF154 [11]) failed in population- or compartment-relevant backgrounds\u2014markers are not wrong in absolute terms, only misassigned; this audit shares the TCGA discovery data and is not independent validation.",
    "disc-audit")

# [58] limitations compression
sub("The premenopausal TCGA subgroup comprised 28 tumours, giving wide positivity confidence intervals (cg24589459 89.3%, 95% CI 72.8\u201396.3%), and external cohorts were not menopause-stratified, so any premenopausal-specific advantage remains to be demonstrated prospectively.",
    "The premenopausal subgroup comprised 28 tumours (cg24589459 positivity 89.3%, 95% CI 72.8\u201396.3%), and external cohorts were not menopause-stratified; any premenopausal-specific advantage remains to be demonstrated prospectively.",
    "lim-premenopausal")
sub("Triage layers used GMQN-normalised \u03b2 values while discovery layers used noob processing; final thresholds will be established in a prospectively collected cohort processed with one predefined laboratory and bioinformatic pipeline.",
    "Discovery and triage layers used different normalisation (noob versus GMQN); final thresholds will be set in a prospectively collected cohort processed with one predefined pipeline.",
    "lim-normalisation")
sub("ZSCAN12 and CELF4 illustrate platform dependence in both directions (absent from EPIC and 450K, respectively), and urine/ctDNA suitability was inferred from blood-cell background without direct validation.",
    "ZSCAN12 and CELF4 illustrate platform dependence (absent from EPIC and 450K respectively); urine/ctDNA suitability was inferred from blood-cell background only.",
    "lim-platform")
sub("The comb-p-like merger does not model inter-probe correlation and may yield optimistic region-level P values; both lead regions were therefore cross-checked with bumphunter (chromosome-wise empirical P<0.002; not a genome-wide FWER).",
    "The comb-p-like merger does not model inter-probe correlation (potentially optimistic region P values), hence the bumphunter cross-check (chromosome-wise empirical P<0.002; not a genome-wide FWER).",
    "lim-combp")
sub("External validation was performed on tumour tissue, not the intended sampling compartments; the benign external controls number only 13; GSE178610 lacks non-malignant controls; and FFPE-derived positivity should not be interpreted as marker performance, because fixation, cohort, tissue-composition and batch differences offer alternative explanations.",
    "External validation used tumour tissue, not the intended sampling compartments; benign external controls number only 13; GSE178610 lacks non-malignant controls; and FFPE positivity is not marker performance\u2014fixation, cohort, tissue-composition and batch differences offer alternative explanations.",
    "lim-extval")

# [5] abstract results
sub("Tumour methylation at panel loci showed no statistically significant menopause-associated differences (28 premenopausal versus 345 postmenopausal tumours), although the analysis was underpowered and equivalence was not established.",
    "No statistically significant menopause-associated methylation differences were detected (28 premenopausal versus 345 postmenopausal tumours), although the analysis was underpowered and equivalence was not established.",
    "abstract-menopause")

# ref 27 access date
sub("Beijing: National Medical Products Administration; 2025. https://english.nmpa.gov.cn/2025-06/11/c_1101553.htm",
    "Beijing: National Medical Products Administration; 2025. https://english.nmpa.gov.cn/2025-06/11/c_1101553.htm (accessed 5 September 2026).",
    "ref27-access-date")

doc.save(F)
w1 = wc()
log.append(f"word count: {w0} -> {w1} ({100*(w1-w0)/w0:+.1f}%)")
print("\n".join(log))
