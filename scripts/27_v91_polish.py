"""v9.1: cross-cohort illustrative OR-rule + 95% CI + abstract wording."""
import math
import re
from pathlib import Path

import numpy as np
from docx import Document

BASE = Path(r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19")
DATA = BASE / "data"
DL = Path(r"C:/Users/ld/Downloads")
SRC = BASE / "Manuscript_EC_methylation_panel_v9.docx"
DST = BASE / "Manuscript_EC_methylation_panel_v9.1.docx"

# ---- ① BOLL OR-rule 跨队列（本地可算部分）----
b1, b2 = "cg24589459", "cg07495363"
need = {b1, b2}

def read_list(fn):
    return [m.group(0) for line in open(DL / fn, encoding="utf-8") for m in [re.search(r"GSM\d+", line)] if m]

def or_stats(gse, ids):
    tp = n = 0
    for g in ids:
        f = DATA / gse / "ewas_betas" / f"{g}.txt"
        if not f.exists():
            continue
        vals = {}
        with open(f) as fh:
            for line in fh:
                pr, b = line.rstrip("\n").split("\t")
                if pr in need and b != "NA":
                    vals[pr] = float(b)
        if b1 in vals and b2 in vals:
            n += 1
            if vals[b1] > 0.3 or vals[b2] > 0.3:
                tp += 1
    return tp, n

ca69 = read_list("NoteGPT_GSE136791_carcinoma_GSM_list.txt")
eec = read_list("NoteGPT_GSE155760_EEC_GSM_list.txt")
ctrl13 = read_list("NoteGPT_GSE155760_endometrial_controls_GSM_list.txt")
s_ca = or_stats("GSE136791", ca69)
s_eec = or_stats("GSE155760", eec)
s_ct = or_stats("GSE155760", ctrl13)
print(f"OR-rule: GSE136791 ca {s_ca[0]}/{s_ca[1]}; GSE155760 EEC {s_eec[0]}/{s_eec[1]}; ctrl {s_ct[0]}/{s_ct[1]}")

# ---- ② Wilson 95% CI ----
def wilson(x, n, z=1.96):
    p = x / n
    den = 1 + z * z / n
    c = p + z * z / (2 * n)
    w = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return (c - w) / den * 100, (c + w) / den * 100

ci_boll = wilson(25, 28)   # 89.3%
ci_zscan = wilson(26, 28)  # 92.9%
print(f"BOLL 89.3% CI {ci_boll[0]:.1f}-{ci_boll[1]:.1f}; ZSCAN12 92.9% CI {ci_zscan[0]:.1f}-{ci_zscan[1]:.1f}")

# ---- 文档修改 ----
doc = Document(SRC)

def set_text(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)

def sub(p, old, new):
    if old in p.text:
        set_text(p, p.text.replace(old, new))
        return True
    return False

hit = {"abstract": False, "ci": False, "rule": False}
for p in doc.paragraphs:
    # ③ 摘要四维表述
    if "Two probes passed the four measurable background dimensions" in p.text:
        hit["abstract"] = sub(p,
            "Two probes passed the four measurable background dimensions (the fifth, symptomatic benign endometrium, could not be directly assessed for cg27577527 because this probe is absent from the EPIC platform)",
            "Both BOLL and ZSCAN12 passed every background dimension on which they were assessable (cg27577527/ZSCAN12 could not be evaluated in the symptomatic-benign EPIC cohort because the probe is absent from the EPIC platform)")
    # ② 局限性补 CI
    if "The premenopausal TCGA subgroup comprised 28 tumours, yielding wide confidence intervals." in p.text:
        hit["ci"] = sub(p,
            "The premenopausal TCGA subgroup comprised 28 tumours, yielding wide confidence intervals.",
            f"The premenopausal TCGA subgroup comprised 28 tumours, so positivity estimates carry wide confidence intervals (cg24589459 89.3%, 95% CI {ci_boll[0]:.1f}–{ci_boll[1]:.1f}%; cg27577527 92.9%, 95% CI {ci_zscan[0]:.1f}–{ci_zscan[1]:.1f}%, Wilson).")
    # ① illustrative rule 跨队列
    if "As an illustrative, non-locked panel rule" in p.text:
        hit["rule"] = sub(p,
            f"yielded 94.2% tumour positivity in the 69 GSE136791 carcinomas and a false-positive rate of 0.0% in the 13 endometrial controls;",
            f"yielded tumour positivity of {s_ca[0]}/{s_ca[1]} (94.2%) in GSE136791 carcinomas and {s_eec[0]}/{s_eec[1]} ({s_eec[0]/s_eec[1]*100:.1f}%) in GSE155760 endometrioid EC, with a false-positive rate of {s_ct[0]}/{s_ct[1]} (0.0%) in the 13 endometrial controls; evaluation in GSE178610 is pending;")

doc.save(DST)
print("saved:", DST)
d2 = Document(DST)
t = "\n".join(p.text for p in d2.paragraphs)
for k, v in hit.items():
    print(k, "OK" if v else "NOT-APPLIED")
print("abstract sentence present:", "every background dimension on which they were assessable" in t)
print("CI present:", "95% CI" in t and "Wilson" in t)
print("cross-cohort rule present:", f"{s_eec[0]}/{s_eec[1]}" in t)
