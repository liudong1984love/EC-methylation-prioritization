"""v9.17 checklist: 11th-round anchors + regression checks."""
import re
from docx import Document

F = r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.17.docx"
d = Document(F)
text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
tt = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
full = text + "\n" + tt

ab = [p.text for p in d.paragraphs if p.text.strip().startswith(("Background:", "Methods:", "Results:", "Conclusions:"))][:4]
aw = sum(len(t.split()) for t in ab)
body = len(text.split())

checks = [
    ("title repositioned", "an in silico evaluation of BOLL and ZSCAN12" in text),
    ("old title gone", "as candidates for premenopausal validation" not in text),
    ("abstract aim toned", "prioritise methylation candidates for future validation in premenopausal EC triage" in text),
    ("abstract proxy removed", "surgical cohort used as an intended-use proxy" not in text),
    ("abstract surgical cohort", "benign surgical cohort than in cancer-free endometrial reference controls" in text),
    ("methods proxy kept once", "used as an approximate intended-use proxy" in text),
    ("primary pipeline declared", "The primary pipeline was prespecified per data source" in text),
    ("old GMQN-triage claim gone", "background triage used GMQN-normalised values" not in text),
    ("TCGA barcode", "matched by TCGA patient barcode" in text),
    ("cycle pairing caveat methods", "best cross-donor r 0.991" in text and "age-concordant" in text),
    ("lit search date x2", text.count("last searched 6 September 2026") == 2),
    ("zscan12 range pipeline clarified", "0.044–0.060 under the primary pipeline of each dimension" in text),
    ("grading sentence", "stable (point-estimate pass under every tested pipeline; both lead probes), pipeline-dependent" in text),
    ("subs GMQN range fixed", "0.158–0.169" in text and "0.160–0.167" not in text),
    ("ERROR 0.025 gone", "0.025" not in full),
    ("ctrl13 anchor 0.059", "P95 β=0.059" in text),
    ("OR rule benign 45.8%", "159/347 (45.8%)" in text),
    ("OR rule endo 33.6%", "214/637 (33.6%)" in text),
    ("OR rule zero cohorts", "0/17" in text and "0/20" in text and "0/42" in text and "0/13" in text),
    ("OR rule no-spec claim", "cannot be read as evidence of high specificity" in text),
    ("tier candidates range", "P95 0.177–0.488" in text),
    ("WID-qEC assertive gone", "absorbs GYPC's probe-level background" not in text),
    ("WID-qEC hedge", "plausible—but untested—interpretation" in text and "cannot be used to evaluate the performance of existing test products" in text),
    ("disc opening toned", "prioritised for future validation in premenopausal AUB triage" in text),
    ("disc hypotheses framing", "generates testable hypotheses about marker choice and sampling design" in text),
    ("conclusion hedge", "generated hypotheses linking the composition of existing tests" in text),
    ("limitations pairing", "pairing is probabilistic" in text),
    ("limitations OR rule", "45.8% rule positivity in the benign surgical cohort" in text),
    ("agenda hyperplasia", "enrolment of both atypical and non-atypical hyperplasia in adequate numbers" in text),
    ("subs independence caveat", "no longer constitute fully independent validation for them" in text),
    ("fig3 legend harmonised", "bar colours use the Table 7 categories" in text),
    ("fig4 legend cross-cohort", "Cross-cohort performance of the BOLL and ZSCAN12 regions" in text),
    ("old fig4 legend gone", "greedy incremental coverage in the TCGA discovery cohort" not in text),
    ("table7 footnote xref", "Figure 3 uses the same colour categories" in text),
    ("table8 anchor cell", "0.059" in tt),
    ("table8 partner cell", "0.266" in tt),
    ("table8 pcdhga cell", "0.352" in tt),
    ("table8 ghrs cell", "0.319" in tt),
    ("table8 arl5c cell", "0.177" in tt),
    ("table8 cdo1 cell", "0.252" in tt),
    ("table8 cyp1b1 cell", "0.488" in tt),
    ("proxy overuse removed", text.count("intended-use-proxy") <= 1),
    ("abstract <=350", aw <= 350),
    ("AH disclaimer kept", "performance for atypical hyperplasia cannot be established" in text),
    ("DOI kept", "10.5281/zenodo.22526552" in text),
    ("no authors", "Dong Liu" not in full and "Yubing Chen" not in full),
    ("PMID count", text.count("PMID:") >= 26),
]

f = 0
for name, ok in checks:
    print(("PASS " if ok else "FAIL ") + name)
    f += not ok
print(f"\nabstract words: {aw} ({[len(t.split()) for t in ab]})")
print(f"body words: {body}")
print(f"{f} FAILURES")
