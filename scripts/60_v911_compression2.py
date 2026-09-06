"""v9.11 compression pass 2: final trims on the two longest paragraphs (47, 58)."""
from docx import Document

SRC = 'Manuscript_EC_methylation_panel_v9.11.docx'

REWRITES = [
("Of the 23 panel probes, 21 are present",
 "Of the 23 panel probes, 21 are present on the EPIC manifest (exceptions cg27577527 [ZSCAN12] and cg20998319); cg27577527 was therefore not evaluable externally. In GSE136791, 13 of 21 evaluable probes reached ≥80% positivity (β>0.3) across all 103 samples and 16 of 21 in the 69 carcinomas, including six at ≥94% and the BOLL anchor cg24589459 at 87.0% (Table 8). In the 23 endometrioid EC samples of GSE155760, positivity was broadly reproducible (91.3–95.7%) but lower for cg24589459 (65.2%), whereas the 13 endometrial controls showed a markedly low background for cg24589459 (P95 β=0.025; 0/13 above β 0.3 for either BOLL probe) compared with cg07495363 (0.270) (Table 8). This dissociation indicates complementary behaviour within the region—cg24589459 low-false-positive, cg07495363 high-sensitivity—and constitutes a post hoc, hypothesis-generating observation made after inspecting these external data; it motivates our proposal to pre-specify and lock a two-probe region-level rule in an independent threshold-setting cohort before any validation. In GSE178610, positivity was preservation-dependent (higher in FFPE than fresh-frozen; Supplementary Data), consistent with fixation- or batch-related inflation; fresh-frozen estimates were used for cross-cohort comparison. cg27577527 was positive in 8/9 tumours (mean β=0.62) in a small 450K cohort not used in discovery (GSE93589). In the 34 GSE136791 hyperplasia samples, positivity remained high (58.8–100%), indicating a neoplasia-gradient rather than carcinoma-specific profile (atypia unannotated; see below). As a post hoc, hypothesis-generating rule derived from these same data (reported for transparency, not as a performance claim), positivity of either BOLL-region probe yielded 65/69 (94.2%) in GSE136791 carcinomas and 22/23 (95.7%) in GSE155760 with 0/13 false positives; we do not report ROC/AUC, because the only cancer-free controls (n=13) come from a different study than the GSE136791 carcinomas and 13 controls are too few for a stable estimate (upper 95% bound ≈23%)."),

("Several limitations warrant emphasis",
 "Several limitations warrant emphasis. The healthy-donor background estimates rest on 17 samples, so P95 statistics are unstable (bootstrap 95% CI for cg07495363 spans 0.069–0.182; per-probe maximum, P90 and proportion above 0.15 in Supplementary Data—for several probes a minority of healthy donors exceed 0.15 even where the P95 passes, so the gate is a screening heuristic rather than a specificity estimate); menopause status was taken from the Xena clinical matrix and was not histologically confirmed. This is an in silico analysis in which selection and coverage evaluation partially share data, so performance estimates are upper bounds. The premenopausal subgroup comprised 28 tumours (cg24589459 positivity 89.3%, 95% CI 72.8–96.3%), external cohorts were not menopause-stratified, and any premenopausal-specific advantage remains to be demonstrated prospectively. Discovery and triage layers used different normalisation (noob versus GMQN); near-gate pass/fail status changed for 11 of 26 probes between pipelines (both lead probes robust), so thresholds are heuristic and pipeline-bound. ZSCAN12 and CELF4 illustrate platform dependence (absent from EPIC and 450K respectively); urine/ctDNA suitability was inferred from blood-cell background only. The comb-p-like merger does not model inter-probe correlation (potentially optimistic region P values), hence the bumphunter cross-check (chromosome-wise empirical P≤0.002, k=0 of B=500 permutations with the +1 correction (k+1)/(B+1)=1/501≈0.002; not a genome-wide FWER). True self-sampling performance additionally depends on tumour-DNA shedding, cellular mixing, DNA degradation and bisulfite conversion efficiency. Tissue-cohort corroboration used tumour tissue, not the intended sampling compartments; benign external controls number only 13; GSE178610 lacks non-malignant controls; and FFPE positivity is not marker performance. GSE223817 lacks per-sample annotation of menopausal status, AUB presentation and benign uterine diagnoses, so this dimension approximates, rather than strictly matches, the intended-use population. No wet-lab validation has been performed; the proposed panel is designed for a premenopausal cohort with matched benign AUB controls."),
]

doc = Document(SRC)
miss = []
done = set()
for p in doc.paragraphs:
    t = p.text.strip()
    for anchor, new in REWRITES:
        if anchor in done:
            continue
        if t.startswith(anchor):
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(new)
            done.add(anchor)
            break
for anchor, _ in REWRITES:
    if anchor not in done:
        miss.append(anchor)
print('MISS:', miss if miss else 'none')
doc.save(SRC)
doc2 = Document(SRC)
words = sum(len(p.text.split()) for p in doc2.paragraphs if p.text.strip())
print('word count now:', words)
