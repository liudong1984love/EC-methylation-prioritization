# -*- coding: utf-8 -*-
"""Chinese manuscript v9.30 -> v9.31 sync (mirrors fix_ms_v931.py V1-V12)."""
import zipfile, re
import xml.etree.ElementTree as ET

SRC = r'C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.30_中文全文翻译.docx'
DST = r'C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.31_中文全文翻译.docx'

def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

zin = zipfile.ZipFile(SRC)
xml = zin.read('word/document.xml').decode('utf-8')

def para_rebuild(xml, marker, new_runs, tag):
    for m in re.finditer(r'<w:p\b[^>]*>.*?</w:p>', xml, re.S):
        block = m.group(0)
        if marker not in block:
            continue
        ppr = re.search(r'<w:pPr>.*?</w:pPr>', block, re.S)
        ppr_s = ppr.group(0) if ppr else ''
        runs = ''.join(
            ('<w:r><w:rPr><w:b w:val="1"/></w:rPr>' if bold else '<w:r>') +
            '<w:t xml:space="preserve">' + esc(text) + '</w:t></w:r>'
            for text, bold in new_runs)
        return xml[:m.start()] + '<w:p>' + ppr_s + runs + '</w:p>' + xml[m.end():]
    raise AssertionError(tag + ': paragraph not found')

REPL = []

# Z1 abstract background
REPL.append(('Z1',
  '检测正快速走向成熟，已有 NMPA 批准的三基因试剂盒、多个国际研究计划和 2026 年中国专家共识。然而，现有标志物主要在绝经后女性中开发',
  '检测正快速走向成熟，但现有标志物主要在绝经后女性中开发'))

# Z2 abstract methods
REPL.append(('Z2',
  '以单探针差异甲基化筛选作为产生候选的主路径；并行的 TCGA 与 GSE67116 差异甲基化区域（DMR）分析仅提供区域背景信息，未向最终池贡献任何候选。候选标志物接受',
  '以单探针差异甲基化筛选作为产生候选的主路径。候选标志物接受'))

# Z3 methods score constants
REPL.append(('Z3',
  '历史实现还附加了两项（跨队列一致性，0.15；扩增子适用性，0.05），它们在该阶段对单探针路径候选为常数，因而不能影响排序；这两项记录在注 S2.3 中，在此不计为证据。',
  '另有两项历史评分项在该阶段对单探针路径候选为常数，因而不能影响排序（注 S2.3）。'))

# Z4 methods donor verified
REPL.append(('Z4',
  '以供者而非组分为单位，供者对应关系由组分内样本顺序推断且未经独立核实——所得区间相近，未改变任何门槛判定',
  '以供者而非组分为单位，供者对应关系已对照 ArrayExpress 样本元数据核实（E-GEOD-35069 样本标题编码供者标识；核实映射见补充数据 S2s）——所得区间与此前按顺序推断的分析完全一致，未改变任何门槛判定'))

# Z5 top-10 caveat
REPL.append(('Z5',
  '门后存活集合通常只有 2–7 个探针，无法支撑前 10 比较；补充数据）',
  '门后存活集合通常只有 2–7 个探针，无法支撑前 10 比较；由于该分析中评分公式固定，排名不变部分源于分析定义本身，而非证明门槛化选择过程的稳定性——门槛化流水线的直接稳定性度量是两个先导探针的条件重抽样保留率（48–53%；见局限性）；补充数据）'))

# Z8 BOLL expression
REPL.append(('Z8',
  '无论甲基化状态如何，BOLL 表达在子宫内膜组织中基本缺失（172 例配对 RNA-seq 的 TCGA 肿瘤；完整分析见补充数据），因此 EC 中的 BOLL 高甲基化最好解读为一个被抑制的生殖系特异位点的异位高甲基化，而非一个表达基因的沉默。',
  '无论甲基化状态如何，BOLL 表达在子宫内膜组织中基本缺失（172 例配对 RNA-seq 的 TCGA 肿瘤与 24 例癌旁正常样本中 log2(x+1) 表达中位数均为 0；表达 >1 的比例肿瘤为 16%、癌旁正常为 0%；甲基化—表达相关性弱负相关，cg24589459 的 Spearman ρ 为 −0.19；完整分析见补充数据）。肿瘤与癌旁子宫内膜组织中 BOLL 表达均低，因此不支持“观察到的甲基化信号源于对一个活跃表达基因的转录沉默”这一解释。'))

# Z9 devoid
REPL.append(('Z9',
  '公共数据仍然缺乏绝经前良性 AUB 的甲基化谱（息肉、肌瘤、子宫腺肌病），这一缺口只有专门建立的队列才能填补。对 EWAS Data Hub 的系统性检索（18 个子宫内膜项目、1,746 例样本）证实了所检索资源内的缺失：我们未找到任何符合条件的公开 450K/EPIC 队列——即提供带所需良性诊断注释的绝经前良性病变（息肉、肌瘤、子宫腺肌病、PCOS）子宫内膜甲基化谱的队列（数据库、检索日期和纳入条件见补充注 S2.11）。',
  '对 EWAS Data Hub 的系统性检索（18 个子宫内膜项目、1,746 例样本）未找到任何符合条件的公开 450K/EPIC 队列——即提供带所需良性诊断注释的绝经前良性病变（息肉、肌瘤、子宫腺肌病、PCOS）子宫内膜甲基化谱的队列（数据库、检索日期和纳入条件见补充注 S2.11）——这一缺口只有专门建立的队列才能填补。'))

# Z10 finalists
REPL.append(('Z10',
  '健康通过集合由 65 减至 41，finalist 由 15 减至 10；全部 5 个 finalist 损失均为临界',
  '健康内膜通过集合由 65 减至 41，第二阶段分流通过集合——位于严格四门 A 级认定上游的辅助阶段，并非 A 级存活者——由 15 减至 10；全部 5 个损失均为临界'))

# Z11 purified-reference
REPL.append(('Z11',
  '这些谱反映纯化参考背景',
  '这些谱反映参考组分背景'))

# Z12 blood fractions donor verified
REPL.append(('Z12',
  '七种纯化细胞类型、全血、PBMC 和粒细胞——均来自相同六名男性供者）[21]',
  '七种纯化细胞类型、全血、PBMC 和粒细胞——均来自相同六名男性供者，供者身份已对照 ArrayExpress 样本标题核实）[21]'))

for tag, old, new in REPL:
    n = xml.count(old)
    assert n == 1, '%s anchor count = %d' % (tag, n)
    xml = xml.replace(old, new, 1)
    print(tag, 'replaced')

# Z6 discovery-threshold paragraph rebuild
xml = para_rebuild(xml, '现阈值敏感性（探索性）。',
  [('发现阈值敏感性（探索性）。 ', True),
   ('由于两个先导探针贴近发现阈值，我们在预设的单侧 27 组合网格上扫描了阈值本身：放宽 Δβ（≥0.50 至 ≥0.30）、总体肿瘤阳性率（≥90% 至 ≥80%）和绝经前阳性率（≥85% 至 ≥75%），FDR、癌旁正常 P95 标准以及经审计的跨队列一致性与注释过滤器保持不变，并对每个网格格点一致施用（补充注 S2.15）。在原始阈值下，全基因组重推导在过滤器之前得到 479 个候选、之后得到 476 个——恰为主分析审计全集——并得到相同的两个 A 级存活者：按有效评分 ZSCAN12 排名 1/2、BOLL 锚定排名 2/2。放宽使过滤后候选集合最多增至 2,942 个探针，但四门背景分流几乎剔除了全部新进入者：在原来的 ≥90% 总体阳性率要求下，任何网格格点都没有新的低背景候选进入；在最松的角点也仅有 2,466 个新进入者中的 17 个通过分流（存活者共 19 个）。cg27577527（ZSCAN12）在全部 27 个组合中均列分流存活者第一；cg24589459（BOLL 锚定）在全部组合中通过分流（按有效评分排名 2–6，变化反映更松阈值下得分更高的低背景新进入者，而非信号丢失）。获批试剂盒与参照标志物呈现不同的排除模式：CDO1 通过全部发现标准，仅在背景分流被剔除；AJAP1 被总体阳性率标准排除（84.2%）；GALR1 和 GYPC 在任何格点均不进入。因此，单探针阳性率低于 90% 并不意味着缺乏互补价值——但放宽阳性率要求也无法挽救试剂盒对照标志物（补充数据 S2l、S2m；跨分析对账见补充注 S2.17，表 S14）。', False)], 'Z6')
print('Z6 rewritten')

# Z7 region-rule paragraph rebuild
xml = para_rebuild(xml, 'OLL 区域水平规则（探索性）。',
  [('BOLL 区域水平规则（探索性）。 ', True),
   ('对五个预先申报的候选规则（方法；补充注 S2.13）按两条预先申报的选择准则进行比较——在良性阳性率 ≤5% 的约束下最大化绝经前肿瘤阳性率（回退：最小化良性阳性率），以及最大化 Youden 指数（绝经前阳性率减良性阳性率）——区域均值规则在两条准则下均为最频繁当选（65.1% 的 bootstrap 重复和 100% 的良性留一折叠；仅锚定 34.9%；OR、AND 和仅配对从未当选），保留 92.9% 的绝经前肿瘤阳性率、良性手术阳性率 2.0%，而 OR 规则为 45.8%（159/347，由配对探针驱动）。肿瘤与良性读数来自不同队列（TCGA 对 GSE223817），重抽样无法消除队列—疾病混杂：这些数字描述的是规则在本数据中的行为，衡量内部规则选择稳定性而非独立验证，其本身并不确立跨队列可迁移性；该规则及其阈值仍有待在独立开发队列中评估并锁定。', False)], 'Z7')
print('Z7 rewritten')

with zipfile.ZipFile(DST, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == 'word/document.xml':
            data = xml.encode('utf-8')
        zout.writestr(item, data)
zin.close()

z = zipfile.ZipFile(DST)
assert z.testzip() is None
ET.fromstring(z.read('word/document.xml'))
x = z.read('word/document.xml').decode('utf-8')
checks = ['检测正快速走向成熟，但现有标志物', '主路径。候选标志物接受',
          '另有两项历史评分项', '对照 ArrayExpress 样本元数据核实',
          '条件重抽样保留率（48–53%', '479 个候选、之后得到 476 个',
          '2,466 个新进入者中的 17 个', '最多增至 2,942 个探针',
          '最大化 Youden 指数', '跨队列可迁移性', '24 例癌旁正常样本',
          '活跃表达基因的转录沉默', '这一缺口只有专门建立的队列才能填补',
          '辅助阶段，并非 A 级存活者', '参考组分背景', '供者身份已对照 ArrayExpress 样本标题核实']
for k in checks:
    print(k[:40], '->', k in x)
print('residual 2,983:', x.count('2,983'), '| 2,507:', x.count('2,507'),
      '| finalist:', x.count('finalist'), '| 异位高甲基化:', x.count('异位高甲基化'),
      '| 未经独立核实:', x.count('未经独立核实'), '| 纯化参考背景:', x.count('纯化参考背景'))
import docx
d = docx.Document(DST)
print('paras:', len([p for p in d.paragraphs if p.text.strip()]), '| tables:', len(d.tables))
print('saved:', DST)
