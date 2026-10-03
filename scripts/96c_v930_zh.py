# -*- coding: utf-8 -*-
"""Chinese full translation v9.29 -> v9.30: mirror of the 13 EN edits (R1-R13)."""
import zipfile, re
import xml.etree.ElementTree as ET

SRC = r'C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.29_中文全文翻译.docx'
DST = r'C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.30_中文全文翻译.docx'

def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

zin = zipfile.ZipFile(SRC)
xml = zin.read('word/document.xml').decode('utf-8')

def para_rebuild(xml, marker, new_runs, tag):
    for m in re.finditer(r'<w:p\b[^>]*>.*?</w:p>', xml, re.S):
        block = m.group(0)
        if marker not in block:
            continue
        ppr = re.search(r'<w:pPr>.*?</w:pPr>', block, re.S)
        ppr_s = ppr.group(0) if ppr else ''
        runs = ''.join(
            ('<w:r><w:rPr><w:b w:val="1"/></w:rPr>' if bold else '<w:r>') +
            '<w:t xml:space="preserve">' + esc(text) + '</w:t></w:r>'
            for text, bold in new_runs)
        return xml[:m.start()] + '<w:p>' + ppr_s + runs + '</w:p>' + xml[m.end():]
    raise AssertionError(tag + ': paragraph not found')

REPL = []

# ---- C2: LOO sentence ----
REPL.append(('C2',
  '且没有任何单一绝经前肿瘤对任一先导探针产生影响（逐一剔除法的阳性率范围分别为 88.9–92.6% 和 92.6–96.3%；补充注 S2.9）',
  '且剔除任何单一绝经前肿瘤均未导致任一先导探针跌破绝经前阳性率标准（逐一剔除法的阳性率范围分别为 88.9–92.6% 和 92.6–96.3%；补充注 S2.9）'))

# ---- C3: Methods rules sentence ----
REPL.append(('C3',
  '并在每次重复中重算发现标准、背景门槛、三项有效评分和候选排名（补充注 S2.9；补充数据 S2f）。',
  '并在每次重复中重算发现标准、背景门槛、三项有效评分和候选排名（补充注 S2.9；补充数据 S2f）。'
  '另外，五个 BOLL 区域水平的事后探索性规则（仅锚定、仅配对、OR、AND 和区域均值；阳性判定 β>0.3）'
  '——在检视外部队列后形成、并在本轮重抽样比较之前预先申报——采用分层 bootstrap（B=2000）和良性样本留一法'
  '进行比较（补充注 S2.13）；选择频率衡量的是内部规则选择稳定性。'))

# ---- C5: Discussion quantitative block -> pointer ----
REPL.append(('C5',
  '为以一个可预设的规则取代这一事后概念，我们比较了五个候选区域水平规则（仅锚定、仅配对、OR、AND 和区域均值；'
  '阳性判定 β>0.3），采用分层 bootstrap（B=2000）和良性样本留一法选择（补充注 S2.13）：'
  '区域均值规则在 65.1% 的 bootstrap 重复和 100% 的良性留一折叠中当选（仅锚定 34.9%；OR、AND 和仅配对从未当选），'
  '保留 92.9% 的绝经前肿瘤阳性率、良性手术阳性率 2.0%，而 OR 规则的良性阳性率为 45.8%（159/347，由配对探针驱动）。据此，',
  '为以一个可预设的规则取代这一事后概念，我们按内部重抽样比较了五个事后探索性候选区域水平规则'
  '（结果；补充注 S2.13），区域均值规则最常当选。据此，'))

# ---- C6: Figure 4 legend ----
REPL.append(('C6',
  '区域均值规则——经 bootstrap 和留一重抽样在 OR、AND 及单探针替代方案中稳定当选（补充注 S2.13）——已被预设用于前瞻性检验。',
  '区域均值规则——在内部 bootstrap 和留一重抽样中最常当选（补充注 S2.13）——已被预设用于前瞻性检验。'))

# ---- C7: upper bounds ----
REPL.append(('C7',
  '选择与覆盖评估部分共用数据，性能估计为上限；',
  '选择与覆盖评估部分共用数据，性能估计可能偏乐观；'))

# ---- C10: no public cohorts ----
REPL.append(('C10',
  '证实了这一缺失：子宫内膜息肉、肌瘤、子宫腺肌病和 PCOS 均无公开的 450K/EPIC 队列（补充注 S2.11）。',
  '证实了所检索资源内的缺失：我们未找到任何符合条件的公开 450K/EPIC 队列——即提供带所需良性诊断注释的'
  '绝经前良性病变（息肉、肌瘤、子宫腺肌病、PCOS）子宫内膜甲基化谱的队列（数据库、检索日期和纳入条件见补充注 S2.11）。'))

# ---- C11: face validity tempered ----
REPL.append(('C11',
  '已知标志物的重新鉴定支持该优选框架的表面效度：在没有任何先验基因清单的情况下，'
  '漏斗重新找到了 WID-qEC 的核心（ZSCAN12 [6,7]）、唯一获批国家试剂盒的锚定基因（CDO1 [9,28]）'
  '和一个荷兰 panel 组分（GHSR [25]）。',
  '已知标志物的重新鉴定支持该优选框架的表面效度，但需作重要限定：在没有任何先验基因清单的情况下，'
  '漏斗经严格背景门槛重新找到了 WID-qEC 的核心（ZSCAN12 [6,7]）；而唯一获批国家试剂盒的锚定基因'
  '（CDO1 [9,28]）和荷兰 panel 组分（GHSR [25]）是作为 B 级对照、凭标本相关背景与既有临床证据保留的，'
  '并非经严格门槛筛出——它们的保留不构成对严格漏斗的独立验证。'))

for tag, old, new in REPL:
    n = xml.count(old)
    assert n == 1, '%s anchor count = %d' % (tag, n)
    xml = xml.replace(old, new, 1)
    print(tag, 'replaced')

# ---- C1: abstract Results compression ----
xml = para_rebuild(xml, '结果：背景估计值在良性手术队列中描述性地高于无癌参考对照。',
  [('结果：背景估计值在良性手术队列中描述性地高于无癌参考对照。cg24589459（BOLL）按点估计通过了全部四道'
    '可评估背景门槛（宫颈门槛为临界：bootstrap 置信区间跨过门槛）；cg27577527（ZSCAN12）通过了探针水平'
    '可评估的三道门槛，背景最为稳定地低（良性手术子宫内膜不可评估：探针不在 EPIC 芯片上）。β>0.3 时绝经前'
    '肿瘤阳性率为 25/28（BOLL）和 26/28（ZSCAN12）；这些估计依赖于选择过程，不是经过验证的灵敏度。'
    '外部组织佐证不完全，简单的双探针 OR 组合在良性手术子宫内膜中频繁阳性（45.8%）——这是灵敏度-背景权衡，'
    '而非已证明的互补性。背景估计依赖于流水线，β 阈值为经验设定，这些候选不是经过验证的 panel。', False)], 'C1')
print('C1 abstract compressed')

# ---- C9: E6 paragraph rewritten (v2 numbers) ----
xml = para_rebuild(xml, '发现阈值敏感性（探索性）。',
  [('发现阈值敏感性（探索性）。 ', True),
   ('由于两个先导探针贴近发现阈值，我们在预设的单侧 27 组合网格上扫描了阈值本身：放宽 Δβ（≥0.50 至 ≥0.30）、'
    '总体肿瘤阳性率（≥90% 至 ≥80%）和绝经前阳性率（≥85% 至 ≥75%），FDR 和癌旁正常 P95 标准保持不变'
    '（补充注 S2.15）。在原始阈值下，全基因组重推导精确复现主分析：479 个候选（476 探针审计全集加 3 个被'
    '审计注释过滤器排除的探针，无一通过分流）、2 个分流存活者，按有效评分 ZSCAN12 排名 1/2、BOLL 锚定排名 2/2。'
    '放宽使候选集合最多增至 2,983 个探针，但四门背景分流几乎剔除了全部新进入者：在原来的 ≥90% 总体阳性率要求下，'
    '任何网格格点都没有新的低背景候选进入；在最松的角点也仅有 2,507 个新进入者中的 17 个通过分流（存活者共 19 个）。'
    'cg27577527（ZSCAN12）在全部 27 个组合中均列分流存活者第一；cg24589459（BOLL 锚定）在全部组合中通过分流'
    '（按有效评分排名 2–6，变化反映更松阈值下得分更高的低背景新进入者，而非信号丢失）。'
    '获批试剂盒与参照标志物呈现不同的排除模式：CDO1 通过全部发现标准，仅在背景分流被剔除；'
    'AJAP1 被总体阳性率标准排除（84.2%）；GALR1 和 GYPC 在任何格点均不进入。'
    '因此，单探针阳性率低于 90% 并不意味着缺乏互补价值——但放宽阳性率要求也无法挽救试剂盒对照标志物'
    '（补充数据 S2l、S2m；跨分析对账见补充注 S2.17，表 S14）。', False)], 'C9')
print('C9 E6 rewritten')

# ---- C8: E7 paragraph rewritten ----
xml = para_rebuild(xml, '按采样部位的并行门控路径（探索性）。',
  [('按采样部位的并行门控路径（探索性）。 ', True),
   ('严格的四门并集筛选的是跨部位低背景标志物，比任何单一预期应用所需都更严格。在一项不改变主规则的并行分析中'
    '（补充注 S2.16），评估了两条与标本匹配的门控路径：子宫内膜刷检路径（门槛：健康和良性手术子宫内膜；'
    '血液作为出血混入风险报告）保留 476 个候选中的 11 个；宫颈/自采样路径（门槛：HPV 阴性宫颈和血液）保留 43 个，'
    '而严格并集下为两个先导探针。这两条路径为实验池提供了有信息量的划分：CDO1——被良性手术维度剔除'
    '（P95 0.372）——通过宫颈路径（宫颈 0.055、血液 0.080），与其已确立的宫颈应用一致；而 6 个池探针呈镜像模式'
    '（子宫内膜通过、宫颈/血液升高；表 S13；补充数据 S2n）。由于良性手术队列缺乏 AUB 注释且与预期使用人群不同，'
    '其背景相对绝经前 AUB 队列的偏差方向和幅度尚不明确；因此宫颈路径的通过不确立子宫内膜适用性。'
    '路径归属指导标本匹配的检测开发；严格四门规则仍为主分析。', False)], 'C8')
print('C8 E7 rewritten')

# ---- C12: future-plan paragraph compressed ----
xml = para_rebuild(xml, '这些局限性定义了一个具体的验证议程',
  [('这些局限性定义的是一个具体的验证议程，而非既定方案。未来工作应优先对核心位点开展靶向评估，'
    '随后进行检测与阈值开发，并在预期使用人群中独立验证；样本量将由预设的准确性与精度目标决定，'
    '设独立的阈值开发集与验证集，并预设入组足够数量的不典型与非不典型增生。'
    '良性对照臂必须覆盖绝经前出血的真实鉴别诊断——子宫内膜息肉、肌瘤、子宫腺肌病和无排卵周期——'
    '而不是健康供者，因为预期使用人群中的背景是特异度的决定性因素。在实验开发上，'
    '我们优先选择临床医生采集的子宫内膜刷样本作为主要标本，宫颈阴道标本保留为次要比较臂；'
    '在此议程中，ZSCAN12 和 CDO1 作为有证据支持的基础候选，BOLL 锚定区域作为探索性附加候选，'
    '决定性问题是 BOLL 能否在同等特异度下找回基础候选漏掉的病例。', False)], 'C12')
print('C12 future plan compressed')

# ---- C4: new Results paragraph after OR-rule paragraph ----
C4 = ('<w:p><w:r><w:rPr><w:b w:val="1"/></w:rPr><w:t xml:space="preserve">'
      + esc('BOLL 区域水平规则（探索性）。 ') + '</w:t></w:r><w:r><w:t xml:space="preserve">'
      + esc('对五个预先申报的候选规则进行比较（方法；补充注 S2.13），区域均值规则在 65.1% 的 bootstrap 重复和 '
            '100% 的良性留一折叠中当选（仅锚定 34.9%；OR、AND 和仅配对从未当选），保留 92.9% 的绝经前肿瘤阳性率、'
            '良性手术阳性率 2.0%，而 OR 规则为 45.8%（159/347，由配对探针驱动）。'
            '这些频率衡量的是内部规则选择稳定性，而非独立验证；该规则及其阈值仍有待在独立开发队列中评估并锁定。')
      + '</w:t></w:r></w:p>')
anchor4 = '任何双探针区域水平规则都必须在包含症状性良性样本的独立阈值设定队列中开发并锁定后，才能进入任何验证。'
n = xml.count(anchor4)
assert n == 1, 'C4 anchor %d' % n
i = xml.find(anchor4)
e = xml.find('</w:p>', i)
xml = xml[:e + len('</w:p>')] + C4 + xml[e + len('</w:p>'):]
print('C4 inserted')

# ---- C13: Table 1 rows (clone GSE223817 row) ----
def clone_row_with_cells(xml, row_marker, new_rows, tag):
    m = re.search(r'<w:tr\b[^>]*>(?:(?!</w:tr>).)*?' + re.escape(row_marker) + r'.*?</w:tr>', xml, re.S)
    assert m, tag + ': row not found'
    block = m.group(0)
    out = ''
    for cells in new_rows:
        nb = block
        ts = list(re.finditer(r'<w:t\b[^>]*>.*?</w:t>', nb, re.S))
        for (txt, tm) in zip(reversed(cells), reversed(ts[:len(cells)])):
            nb = nb[:tm.start()] + '<w:t xml:space="preserve">' + esc(txt) + '</w:t>' + nb[tm.end():]
        out += nb
    return xml[:m.end()] + out + xml[m.end():]

xml = clone_row_with_cells(xml, 'GSE223817', [
  ['GSE87621', '450K', '5 例良性子宫内膜样本', '背景：支持性良性（ZSCAN12 首次直接评估；S2.11）'],
  ['GSE73948', '450K', '26 例内异症患者在位子宫内膜（分泌期）', '背景：支持性良性（S2.11）'],
], 'C13')
print('C13 Table 1 rows added')

with zipfile.ZipFile(DST, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == 'word/document.xml':
            data = xml.encode('utf-8')
        zout.writestr(item, data)
zin.close()

z = zipfile.ZipFile(DST)
assert z.testzip() is None
ET.fromstring(z.read('word/document.xml'))
x = z.read('word/document.xml').decode('utf-8')
checks = ['跌破绝经前阳性率标准', '事后探索性规则', '最常当选', '可能偏乐观',
          '未找到任何符合条件的公开 450K/EPIC 队列', '不构成对严格漏斗的独立验证',
          '而非既定方案', '精确复现主分析', '排名 2–6', '偏差方向和幅度尚不明确',
          'BOLL 区域水平规则（探索性）。', 'GSE87621', 'GSE73948', '补充注 S2.17']
for k in checks:
    print(k[:40], '->', k in x)
print('residual 为上限:', x.count('性能估计为上限'))
print('residual 很可能高估:', x.count('很可能高估'))
print('residual 稳定当选:', x.count('稳定当选'))
print('residual 产生影响:', x.count('对任一先导探针产生影响'))
import docx
d = docx.Document(DST)
print('tables:', len(d.tables), '| table1 rows:', len(d.tables[0].rows))
print('saved:', DST)
