"""v9.17 -> v9.18: 12th review round.
Trade-off reframing (not complementarity); unpaired cycle comparison as primary;
threshold-scan basis clarified; fold range corrected (1.0-1.7, median 1.4);
product-design explanations removed; citation fixes (CELF4->[8]; trend ref [29]);
wording (passenger, clinically verified, prespecified audit)."""
from docx import Document

SRC = r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.17.docx"
DST = r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.18.docx"

doc = Document(SRC)
log = []

def set_text(p, new):
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""

def sub(p, old, new, tag):
    if old in p.text:
        set_text(p, p.text.replace(old, new))
        log.append("OK   " + tag)
        return True
    log.append("MISS " + tag)
    return False

# ---------------- Abstract ----------------
for p in doc.paragraphs:
    t = p.text
    if t.startswith("Results:"):
        sub(p, "met the prespecified gate by point estimate",
            "met the primary gate by point estimate", "abstract prespecified")
        sub(p, "Corroboration in independent tumour-tissue cohorts—not the intended sampling compartments—was partial and yielded a post hoc two-probe region-level hypothesis for BOLL.",
            "Corroboration in independent tumour-tissue cohorts—not the intended sampling compartments—was partial; the two BOLL-region probes showed a sensitivity–background trade-off, not demonstrated complementarity—a post hoc two-probe OR rule was frequently positive in benign surgical endometrium.", "abstract trade-off")
        sub(p, "although the analysis was underpowered and did not establish equivalence",
            "although underpowered to establish equivalence", "abstract trim1")
        sub(p, "Discovery-cohort positivity and coverage estimates are optimistic upper bounds",
            "Discovery-cohort estimates are optimistic upper bounds", "abstract trim2")
        sub(p, "(probe absent from the EPIC platform)",
            "(probe absent from EPIC)", "abstract trim3")
    if t.startswith("Methods:"):
        sub(p, "with harmonised differential-methylation criteria and region merging",
            "with harmonised criteria and region merging", "abstract trim4")
        sub(p, "were audited through the same funnel",
            "were audited identically", "abstract trim5")
        sub(p, "five background dimensions spanning healthy endometrium",
            "five background dimensions: healthy endometrium", "abstract trim6")

# ---------------- Introduction ----------------
for p in doc.paragraphs:
    if p.text.startswith("Endometrial carcinoma (EC) is the most common"):
        sub(p, "with a concerning shift toward younger, premenopausal women [1].",
            "with a concerning shift toward younger, premenopausal women [1,29].", "intro trend ref")

# ---------------- Methods ----------------
for p in doc.paragraphs:
    t = p.text
    if t.startswith("Candidates were triaged against five dimensions"):
        sub(p, "(iii) early/mid-secretory endometrium (GSE90060; the source publication reports paired donor sampling [13], but donor identifiers are not annotated in the accessed records, so 17 pairs were reconstructed by genome-wide methylation similarity—within-pair r median 0.992, best cross-donor r 0.991, 15/17 pairs age-concordant—and drift estimates were confirmed on the age-concordant subset);",
            "(iii) early/mid-secretory endometrium (GSE90060; donor identifiers are not annotated in the accessed records, so the primary analysis compares early- and mid-secretory groups unpaired; the source publication reports paired donor sampling [13], and a sensitivity analysis on 17 methylation-similarity-inferred pairs—within-pair r median 0.992, best cross-donor r 0.991, 15/17 age-concordant—gave concordant estimates);", "methods cycle unpaired")
    if t.startswith("Analyses used R 4.6.0"):
        sub(p, "The primary pipeline was prespecified per data source",
            "The primary pipeline was defined per data source", "methods prespecified")
        sub(p, "Candidate retention was scanned across background P95 gates of 0.05–0.30.",
            "Candidate retention was scanned across background P95 gates of 0.05–0.30, using for each probe and dimension the least favourable estimate across pipelines (a conservative choice).", "methods scan basis")
        sub(p, "the GMQN values used in triage", "the GMQN values distributed for these cohorts", "methods gmqn triage 1")
        sub(p, "the GMQN values used for triage", "the GMQN values distributed for these cohorts", "methods gmqn triage 2")

# ---------------- Results ----------------
for p in doc.paragraphs:
    t = p.text
    if t.startswith("All five markers of the internally derived set"):
        sub(p, "LOC134466 and HAND2 (blood-cell backgrounds 0.603 and 0.371; largest menstrual-cycle drift)",
            "LOC134466 and HAND2 (blood-cell backgrounds 0.603 and 0.371; cycle-phase group-mean differences 0.026–0.037)", "internal audit drift")
    if t.startswith("Background P95 β values in the benign surgical"):
        sub(p, "approximately 1.2- to 1.6-fold",
            "1.0- to 1.7-fold across the tiered candidates (median ≈1.4)", "fold range fix")
    if t.startswith("The two lead probes differ in evidentiary status"):
        sub(p, "Menstrual-cycle drift of the anchor probe was minimal (mean |Δβ| 0.026 between fingerprint-inferred early/mid-secretory pairs, 0.028 on the 15 age-concordant pairs; experimental-pool range 0.014–0.042, concordant-subset 0.013–0.044; pairing is inference—see Methods); proliferative-phase, menstrual and anovulatory states were not assayed.",
            "Between early- and mid-secretory endometrium the anchor probe showed a small unpaired group-mean difference (|Δβ|=0.005; experimental-pool range <0.001–0.020; a sensitivity analysis on methylation-similarity-inferred pairs gave concordant estimates, mean |Δβ| 0.026—see Methods); only two secretory phases were assayed—proliferative-phase, menstrual and anovulatory states were not covered—so robustness to the full menstrual cycle cannot be claimed.", "cycle drift unpaired main")
    if t.startswith("Auditing the NMPA-kit and WID-qEC components"):
        sub(p, "CELF4, a principal marker of the Taiwanese series [8,25], has no 450K probe and was invisible to our pipeline.",
            "CELF4, the partner marker of CDO1 in the CISENDO assay [8], has no 450K probe and was invisible to our pipeline.", "celf4 citation fix")
    if t.startswith("Three candidates showed favourable"):
        sub(p, "did not meet our prespecified probe-level background criteria for cervical or self-sampling applications because of blood-cell background, consistent with their deployment via an intrauterine brush [10];",
            "did not meet our probe-level background criteria for cervical or self-sampling applications because of blood-cell background;", "compartment prespecified+brush")
    if t.startswith("Threshold sensitivity analysis distinguished"):
        sub(p, "cg27577527 passed every assessable dimension down to a gate of 0.09, whereas",
            "cg27577527 passed every assessable dimension down to a gate of 0.09—its entry point reflecting the least-favourable cross-pipeline estimates used in the scan (GMQN healthy-endometrium P95 0.085 versus primary noob 0.060)—whereas", "scan 0.09 explained")
    if t.startswith("Cross-pipeline consistency."):
        sub(p, "only the stable grade supports pipeline-independent candidate claims",
            "only the stable grade supports candidate claims robust within the tested pipelines and datasets", "stable grade wording")
    if t.startswith("Of the 23 experimental-pool probes, 21 are present"):
        sub(p, "This dissociation indicates complementary behaviour within the region—cg24589459 low-background, cg07495363 high-sensitivity—and constitutes a post hoc, hypothesis-generating observation made after inspecting these external data.",
            "This dissociation constitutes a sensitivity–background trade-off rather than demonstrated diagnostic complementarity—cg24589459 low-background but less consistently positive externally, cg07495363 high-sensitivity but with appreciable benign background—and is a post hoc, hypothesis-generating observation made after inspecting these external data.", "p49 trade-off")

# ---------------- Discussion ----------------
for p in doc.paragraphs:
    t = p.text
    if t.startswith("Methylation-based EC detection has entered clinical reality"):
        sub(p, "and a reproducible audit framework that generates testable hypotheses about marker choice and sampling design.",
            "and a reproducible audit framework that quantifies how candidate background varies with cohort and sampling compartment.", "disc opening 3")
    if t.startswith("Re-identification of established markers supports"):
        sub(p, "Other components did not meet our prespecified probe-level background criteria",
            "Other components did not meet our probe-level background criteria", "disc prespecified")
        sub(p, "These probe-level findings do not evaluate the exact CpGs, amplicons or classification algorithms used in the commercial or published assays. One plausible—but untested—interpretation is that region-level quantification (such as ΣPMR) tolerates probe-level background that eliminates single probes, and that intrauterine-brush deployment bypasses blood-cell-dominated compartments; we did not assay the actual amplicons or classifiers, and these probe-level observations cannot be used to evaluate the performance of existing test products.",
            "These probe-level findings do not evaluate the exact CpGs, amplicons or classification algorithms used in the commercial or published assays; they motivate assay-specific evaluation of background in the intended specimen type, and they do not explain or assess the performance of existing assays.", "product explanation removed")
    if t.startswith("No EC-detection report for BOLL"):
        sub(p, "and the expression analysis marks BOLL as a passenger rather than driver marker, bounding biological interpretation without reducing diagnostic utility.",
            "and the expression analysis argues against interpreting the methylation signal as silencing of an originally expressed gene, bounding biological interpretation without reducing diagnostic utility.", "passenger wording")
    if t.startswith("Several limitations warrant emphasis"):
        sub(p, "menopause status was taken from the Xena clinical matrix and was not histologically confirmed",
            "menopausal status was taken from the Xena clinical matrix and was not independently clinically verified", "menopause wording")
        sub(p, "Menstrual-cycle pairing in GSE90060 was inferred from genome-wide methylation similarity because donor identifiers are not annotated; within-pair and best cross-donor correlations are close (median 0.992 versus 0.991), so pairing is probabilistic, although drift estimates changed little on the 15 age-concordant pairs.",
            "GSE90060 donor identifiers are not annotated, so the primary cycle comparison is an unpaired group contrast covering only two secretory phases; pairing by methylation similarity is probabilistic (within-pair versus best cross-donor r median 0.992 versus 0.991) and is reported only as a sensitivity analysis.", "limitations pairing reframe")
    if t.startswith("In conclusion, multi-cohort discovery"):
        sub(p, "—and generated hypotheses linking the composition of existing tests to their sampling designs.",
            ".", "conclusion clause removed")
        if "generated hypotheses linking the composition" in p.text:
            sub(p, " and generated hypotheses linking the composition of existing tests to their sampling designs", "", "conclusion clause removed 2")

# ---------------- Figure legend ----------------
for p in doc.paragraphs:
    if p.text.startswith("Figure 4."):
        sub(p, "illustrate the complementary behaviour motivating region-level assay design.",
            "illustrate the sensitivity–background trade-off that region-level assay design must resolve.", "fig4 legend trade-off")

# ---------------- Table 4 drift column -> unpaired ----------------
DRIFT = {"cg24589459": "0.005", "cg27577527": "0.004", "cg15790037": "0.008",
         "cg16439198": "<0.001", "cg23180938": "0.018", "cg18507379": "0.012", "cg10109500": "0.005"}
t4 = doc.tables[3]
fixed = 0
for row in t4.rows:
    if row.cells[1].text.strip() in DRIFT:
        set_text(row.cells[6].paragraphs[0], DRIFT[row.cells[1].text.strip()])
        fixed += 1
log.append(f"OK   table4 drift cells: {fixed}/7")

# Table 4 note (cycle drift definition)
for p in doc.paragraphs:
    if p.text.startswith("Cycle drift: mean |Δβ| between fingerprint-paired"):
        set_text(p, "Cycle drift: absolute group-mean difference between early- and mid-secretory endometrium (GSE90060, unpaired; a sensitivity analysis on methylation-similarity-inferred pairs gave concordant estimates—Supplementary Data); smaller values indicate less phase-associated difference (two secretory phases only).")
        log.append("OK   table4 drift note")

# ---------------- Reference [29] ----------------
REF29 = ("Liu L, Habeshian TS, Zhang J, Peeri NC, Du M, De Vivo I, et al. Differential trends in rising "
         "endometrial cancer incidence by age, race, and ethnicity. JNCI Cancer Spectr. 2023;7(1):pkad001. "
         "doi:10.1093/jncics/pkad001. PMID:36625534.")
ref_done = False
paras = doc.paragraphs
for i, p in enumerate(paras):
    if p.text.startswith("Zhou W, Laird PW, Shen H."):
        if i + 1 < len(paras):
            np_ = paras[i + 1].insert_paragraph_before(REF29)
        else:
            np_ = doc.add_paragraph(REF29)
        if p.runs:
            np_.style = p.style
        ref_done = True
        break
log.append("OK   ref [29] added" if ref_done else "MISS ref [29]")

doc.save(DST)
print("saved:", DST)
for l in log:
    print(l)
print("MISSES:", sum(1 for l in log if l.startswith("MISS")))

# word count check
d2 = Document(DST)
for p in d2.paragraphs:
    if p.text.startswith(("Background:", "Methods:", "Results:", "Conclusions:")):
        print(p.text[:12], "words:", len(p.text.split()))
