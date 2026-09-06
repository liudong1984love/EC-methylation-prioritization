"""v9.3: register GSE93589, fix false-independence, abstract clarification, OR-rule AUC."""
import re
from pathlib import Path

import numpy as np
from docx import Document

BASE = Path(r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19")
DATA = BASE / "data"
DL = Path(r"C:/Users/ld/Downloads")
SRC = BASE / "Manuscript_EC_methylation_panel_v9.2.docx"
DST = BASE / "Manuscript_EC_methylation_panel_v9.3.docx"

b1, b2 = "cg24589459", "cg07495363"
need = {b1, b2}

def read_list(fn):
    return [m.group(0) for l in open(DL / fn, encoding="utf-8") for m in [re.search(r"GSM\d+", l)] if m]

def scores(gse, ids):
    out = []
    for g in ids:
        f = DATA / gse / "ewas_betas" / f"{g}.txt"
        if not f.exists():
            continue
        vals = {}
        with open(f) as fh:
            for line in fh:
                pr, b = line.rstrip("\n").split("\t")
                if pr in need and b != "NA":
                    vals[pr] = float(b)
        if b1 in vals and b2 in vals:
            out.append(max(vals[b1], vals[b2]))
    return np.array(out)

def auc(pos, neg):
    # Mann-Whitney AUC
    allv = np.concatenate([pos, neg])
    ranks = np.argsort(np.argsort(allv))
    rp = ranks[: len(pos)].sum() + len(pos)  # ranks are 0-based
    return (rp - len(pos) * (len(pos) + 1) / 2) / (len(pos) * len(neg))

ca69 = read_list("NoteGPT_GSE136791_carcinoma_GSM_list.txt")
eec = read_list("NoteGPT_GSE155760_EEC_GSM_list.txt")
ctrl13 = read_list("NoteGPT_GSE155760_endometrial_controls_GSM_list.txt")
s_ca, s_eec, s_ct = scores("GSE136791", ca69), scores("GSE155760", eec), scores("GSE155760", ctrl13)
a1 = auc(s_ca, s_ct)
a2 = auc(s_eec, s_ct)
print(f"OR-score AUC: GSE136791 ca69 vs ctrl13 = {a1:.3f}; GSE155760 EEC vs ctrl13 = {a2:.3f}")

doc = Document(SRC)

def set_text(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)

def try_all(old, new, tag):
    for p in doc.paragraphs:
        if old in p.text:
            set_text(p, p.text.replace(old, new))
            print("OK  ", tag)
            return True
    print("MISS", tag)
    return False

# 1) 假独立修正
try_all(
    "Independently, the 8 hyperplasia samples of GSE67116 (450K) also showed elevated background at these loci (P95 β 0.44–0.69), corroborating the neoplasia-gradient pattern observed in GSE136791.",
    "Consistently with this pattern—though not constituting independent evidence, as GSE67116 contributed to discovery—the 8 hyperplasia samples of GSE67116 (450K) also showed elevated background at these loci (P95 β 0.44–0.69).",
    "false-independence")

# 2) 摘要分解句写明比较两端
try_all(
    "cancer-free vs symptomatic endometrial controls: P95 β 0.059 vs 0.100 for cg24589459; 0.266 vs 0.427 for cg07495363; 0.252 vs 0.372 for CDO1",
    "cancer-free endometrial controls (GSE155760, n=13) vs symptomatic controls (GSE223817, n=347), both on EPIC/GMQN: P95 β 0.059 vs 0.100 for cg24589459; 0.266 vs 0.427 for cg07495363; 0.252 vs 0.372 for CDO1",
    "abstract-ends")

# 3) Methods 数据源补 GSE93589
try_all(
    "GSE90060, GSE46306, GSE35069 and GSE223817 were obtained as GMQN-normalised β values with curated metadata from the NGDC EWAS Data Hub [16].",
    "GSE90060, GSE46306, GSE35069, GSE223817, GSE136791, GSE155760 and GSE93589 were obtained as GMQN-normalised β values with curated metadata from the NGDC EWAS Data Hub [16].",
    "methods-gse93589")

# 4) OR 规则句追加 AUC
try_all(
    "evaluation in GSE178610 is pending;",
    f"evaluation in GSE178610 is pending; as an illustrative discrimination estimate, the two-probe BOLL-region score (maximum of the two probes) separated carcinomas from the 13 endometrial controls with an AUC of {a1:.2f} (GSE136791) and {a2:.2f} (GSE155760);",
    "or-auc")

doc.save(DST)

# 5) Table 1 加 GSE93589 行
doc = Document(DST)
t1 = doc.tables[0]
row = t1.add_row()
vals = ["GSE93589", "450K", "9 endometrial carcinomas", "External validation (ZSCAN12/BOLL, 450K-only probes)"]
for c, v in zip(row.cells, vals):
    c.text = ""
    c.paragraphs[0].add_run(v)
doc.save(DST)
print("saved:", DST)
print("table1 rows:", len(doc.tables[0].rows))
