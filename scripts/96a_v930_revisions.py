# -*- coding: utf-8 -*-
"""Manuscript v9.29 -> v9.30: numeric reconciliation + pre-submission compression.
R1  Abstract Results compression.
R2  LOO sentence: 'influenced either lead' -> 'did not cause either lead to fail'.
R3  Methods: post hoc exploratory BOLL region rules (pre-declared before resampling).
R4  Results: new short paragraph reporting the five-rule comparison.
R5  Discussion: quantitative block -> pointer (results now in Results section).
R6  Figure 4 legend: 'stably selected' -> 'most frequently selected in internal resampling'.
R7  Limitations: 'upper bounds' -> 'may be optimistically biased'.
R8  E7 paragraph: remove AUB-overestimate/conservative inference (direction unknown).
R9  E6 paragraph: rewritten with C4 v2 numbers (original cell reproduces main analysis).
R10 'no public cohorts exist' -> 'we identified no eligible ... in the resources searched'.
R11 Face-validity sentence tempered (Tier-B retention != strict-funnel validation).
R12 Future-plan paragraph compressed (drop 50+100 / 140+200 numbers).
R13 Table 1: add GSE87621 and GSE73948 rows.
"""
import zipfile, re
import xml.etree.ElementTree as ET

SRC = r'C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.29.docx'
DST = r'C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.30.docx'

B = '\u03b2'; LE = '\u2264'; GE = '\u2265'; MD = '\u2014'; DQ = '\u2019'; DL = '\u0394'

def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

zin = zipfile.ZipFile(SRC)
xml = zin.read('word/document.xml').decode('utf-8')

def para_rebuild(xml, marker, new_runs, tag, must_start=None):
    """Replace the whole <w:p> containing marker with new runs [(text,bold),...]."""
    for m in re.finditer(r'<w:p\b[^>]*>.*?</w:p>', xml, re.S):
        block = m.group(0)
        if marker not in block:
            continue
        ppr = re.search(r'<w:pPr>.*?</w:pPr>', block, re.S)
        ppr_s = ppr.group(0) if ppr else ''
        runs = ''.join(
            ('<w:r><w:rPr><w:b w:val="1"/></w:rPr>' if bold else '<w:r>') +
            '<w:t xml:space="preserve">' + esc(text) + '</w:t></w:r>'
            for text, bold in new_runs)
        return xml[:m.start()] + '<w:p>' + ppr_s + runs + '</w:p>' + xml[m.end():]
    raise AssertionError(tag + ': paragraph not found')

REPL = []

# ---- R2: LOO sentence ----
REPL.append(('R2',
  'and no single premenopausal tumour influenced either lead (leave-one-out positivity ranges '
  '88.9\u201392.6% and 92.6\u201396.3%; Supplementary Note S2.9)',
  'and removing any single premenopausal tumour did not cause either lead to fail the '
  'premenopausal-positivity criterion (leave-one-out positivity ranges 88.9\u201392.6% and '
  '92.6\u201396.3%; Supplementary Note S2.9)'))

# ---- R3: Methods, post hoc exploratory rules ----
REPL.append(('R3',
  'recomputed per replicate (Supplementary Note S2.9; Supplementary Data S2f).',
  'recomputed per replicate (Supplementary Note S2.9; Supplementary Data S2f). Separately, five '
  'post hoc exploratory region-level rules for BOLL (anchor only, partner only, OR, AND and '
  'region mean; positivity ' + B + '>0.3), formulated after inspection of the external cohorts '
  'and pre-declared before their resampling comparison, were compared by stratified bootstrap '
  '(B=2000) and benign leave-one-out selection (Supplementary Note S2.13); the selection '
  'frequencies measure internal rule-selection stability.'))

# ---- R5: Discussion quantitative block -> pointer ----
REPL.append(('R5',
  'To replace this post-hoc concept with a pre-specifiable rule, five candidate region-level '
  'rules (anchor only, partner only, OR, AND and region mean; positivity ' + B + '>0.3) were '
  'compared by stratified bootstrap (B=2000) and benign leave-one-out selection (Supplementary '
  'Note S2.13): the region-mean rule was selected in 65.1% of bootstrap replicates and 100% of '
  'benign leave-one-out folds (anchor-only 34.9%; OR, AND and partner-only never selected), '
  'retaining 92.9% premenopausal tumour positivity with 2.0% benign surgical positivity, versus '
  '45.8% benign positivity under the OR rule (159/347, partner-driven). Accordingly,',
  'To replace this post-hoc concept with a pre-specifiable rule, five post hoc exploratory '
  'candidate region-level rules were compared by internal resampling (Results; Supplementary '
  'Note S2.13), which most frequently favoured the region-mean rule. Accordingly,'))

# ---- R6: Figure 4 legend ----
REPL.append(('R6',
  'the region-mean rule' + MD + 'stably selected over OR, AND and single-probe alternatives by '
  'bootstrap and leave-one-out resampling (Supplementary Note S2.13)' + MD + 'is pre-specified '
  'for prospective testing.',
  'the region-mean rule' + MD + 'most frequently selected over OR, AND and single-probe '
  'alternatives in internal bootstrap and leave-one-out resampling (Supplementary Note '
  'S2.13)' + MD + 'is pre-specified for prospective testing.'))

# ---- R7: upper bounds ----
REPL.append(('R7',
  'so performance estimates are upper bounds',
  'so performance estimates may be optimistically biased'))

# ---- R10: no public cohorts ----
REPL.append(('R10',
  'confirmed this absence: no public 450K/EPIC cohorts exist for endometrial polyps, leiomyomas, '
  'adenomyosis or PCOS (Supplementary Note S2.11).',
  'confirmed this absence within the resources searched: we identified no eligible publicly '
  'available 450K/EPIC cohort providing premenopausal benign-lesion (polyp, leiomyoma, '
  'adenomyosis or PCOS) endometrial methylation profiles with the required benign-diagnosis '
  'annotations (databases, search date and eligibility criteria in Supplementary Note S2.11).'))

# ---- R11: face validity tempered ----
REPL.append(('R11',
  'Re-identification of established markers supports the face validity of the prioritisation '
  'framework: without any prior gene list, the funnel re-identified the core of WID-qEC '
  '(ZSCAN12 [6,7]), the anchor of the only approved national kit (CDO1 [9,28]) and a '
  'Dutch-panel component (GHSR [25]).',
  'Re-identification of established markers supports the face validity of the prioritisation '
  'framework, with an important qualification: without any prior gene list, the funnel '
  're-identified the core of WID-qEC (ZSCAN12 [6,7]) through the strict background gates, '
  'whereas the anchor of the only approved national kit (CDO1 [9,28]) and a Dutch-panel '
  'component (GHSR [25]) were retained as Tier-B comparators on specimen-relevant background '
  'and prior clinical evidence, not by the strict gates; their retention therefore does not '
  'constitute independent validation of the strict funnel.'))

for tag, old, new in REPL:
    n = xml.count(old)
    assert n == 1, '%s anchor count = %d' % (tag, n)
    xml = xml.replace(old, new, 1)
    print(tag, 'replaced')

# ---- R1: Abstract Results compression (whole paragraph) ----
xml = para_rebuild(xml, 'Background estimates were descriptively higher in the benign surgical cohort than in cancer-free reference controls.',
  [('Results: Background estimates were descriptively higher in the benign surgical cohort than '
    'in cancer-free reference controls. cg24589459 (BOLL) met all four assessable background '
    'gates by point estimate (cervical gate borderline: bootstrap CI crosses it); cg27577527 '
    '(ZSCAN12) met the three probe-assessable gates with the most consistently low background '
    '(benign surgical endometrium not evaluable: probe absent from EPIC). Premenopausal tumour '
    'positivity was 25/28 (BOLL) and 26/28 (ZSCAN12) at ' + B + '>0.3; these estimates are '
    'selection-dependent, not validated sensitivities. External tissue corroboration was '
    'partial, and a simple two-probe OR combination was frequently positive (45.8%) in benign '
    'surgical endometrium' + MD + 'a sensitivity' + MD + 'background trade-off, not demonstrated '
    'complementarity. Background estimates were pipeline-dependent, so ' + B + ' thresholds are '
    'heuristic, and the candidates are not a validated panel.', False)], 'R1')
print('R1 abstract compressed')

# ---- R9: E6 paragraph rewritten with C4 v2 numbers ----
xml = para_rebuild(xml, 'Discovery-threshold sensitivity (exploratory).',
  [('Discovery-threshold sensitivity (exploratory). ', True),
   ('Because the leads sit close to the discovery cut-offs, the cut-offs themselves were scanned '
    'over a pre-specified, one-sided 27-combination grid loosening ' + DL + B + ' (' + GE +
    '0.50 to ' + GE + '0.30), overall tumour positivity (' + GE + '90% to ' + GE + '80%) and '
    'premenopausal positivity (' + GE + '85% to ' + GE + '75%), with FDR and the '
    'adjacent-normal P95 criterion held fixed (Supplementary Note S2.15). At the original '
    'cut-offs the genome-wide re-derivation reproduces the main analysis exactly: 479 '
    'candidates (the 476-probe audit universe plus three probes excluded by the audited '
    'annotation filters, none passing triage), two triage survivors, ZSCAN12 rank 1/2 and BOLL '
    'anchor rank 2/2 by the effective score. Relaxation enlarged the candidate set to up to '
    '2,983 probes, but the four-gate background triage eliminated nearly all entrants: under '
    'the original ' + GE + '90% overall-positivity requirement no new low-background candidate '
    'entered in any grid cell, and at the loosest corner only 17 of 2,507 new entrants passed '
    'triage (19 survivors in total). cg27577527 (ZSCAN12) ranked first among triage survivors '
    'in all 27 combinations; cg24589459 (BOLL anchor) survived triage in all combinations '
    '(rank 2' + MD + '6 by the effective score, the variation reflecting higher-scoring '
    'low-background entrants at looser thresholds rather than signal loss). Approved-kit and '
    'reference markers showed distinct exclusion modes: CDO1 passes all discovery criteria and '
    'is eliminated only at background triage; AJAP1 is excluded by the overall-positivity '
    'criterion (84.2%); GALR1 and GYPC enter at no grid cell. Single-probe positivity below '
    '90% therefore does not imply absence of complementary value' + MD + 'but relaxing the '
    'positivity requirement does not rescue the kit comparators either (Supplementary Data '
    'S2l, S2m; cross-analysis reconciliation in Supplementary Note S2.17, Table S14).', False)], 'R9')
print('R9 E6 rewritten')

# ---- R8: E7 paragraph rewritten (conservative inference removed) ----
xml = para_rebuild(xml, 'Application-specific gating paths (exploratory).',
  [('Application-specific gating paths (exploratory). ', True),
   ('The strict four-gate union selects cross-compartment low-background markers and is more '
    'stringent than any single intended application requires. In a parallel analysis that does '
    'not alter the main rule (Supplementary Note S2.16), two specimen-matched gating paths '
    'were evaluated: an endometrial-brush path (gates: healthy and benign-surgical '
    'endometrium; blood reported as bleeding-contamination risk) retained 11 of 476 '
    'candidates, and a cervical/self-sampling path (gates: HPV-negative cervix and blood) '
    'retained 43, versus the two leads under the strict union. The paths partition the '
    'experimental pool informatively: CDO1' + MD + 'eliminated by the benign-surgical '
    'dimension (P95 0.372)' + MD + 'passes the cervical path (cervix 0.055, blood 0.080), '
    'consistent with its established cervical application, whereas six pool probes show the '
    'mirror pattern (endometrium-passing, cervix/blood-elevated; Table S13; Supplementary '
    'Data S2n). Because the benign-surgical cohort lacks AUB annotation and differs from the '
    'intended-use population, the direction and magnitude of background bias relative to a '
    'premenopausal AUB cohort remain unknown; cervical-path passes therefore do not establish '
    'endometrial suitability. Path assignments guide specimen-matched assay development; the '
    'strict four-gate rule remains the main analysis.', False)], 'R8')
print('R8 E7 rewritten')

# ---- R12: future-plan paragraph compressed ----
xml = para_rebuild(xml, 'These limitations define a concrete validation agenda',
  [('These limitations define a concrete validation agenda rather than a settled protocol. '
    'Future work should prioritise targeted assessment of the core loci, followed by assay and '
    'threshold development and independent validation in the intended-use population; sample '
    'size will be determined by prespecified accuracy and precision targets, with separate '
    'threshold-development and validation sets and prespecified enrolment of atypical and '
    'non-atypical hyperplasia. The benign comparator arm must span the realistic differential '
    'diagnosis of premenopausal bleeding' + MD + 'endometrial polyps, leiomyomas, adenomyosis '
    'and anovulatory cycles' + MD + 'rather than healthy donors, because background in the '
    'intended-use population is the decisive determinant of specificity. For experimental '
    'development we prioritise clinician-collected endometrial brush samples as the primary '
    'specimen, with cervicovaginal specimens retained as a secondary comparison arm; within '
    'this agenda ZSCAN12 and CDO1 serve as evidence-backed base candidates and the BOLL '
    'anchor region as an exploratory add-on, the decisive question being whether BOLL '
    'recovers, at equal specificity, cases missed by the base candidates.', False)], 'R12')
print('R12 future plan compressed')

# ---- R4: new Results paragraph after the OR-rule paragraph ----
R4 = ('<w:p><w:r><w:rPr><w:b w:val="1"/></w:rPr><w:t xml:space="preserve">'
      + esc('Region-level rules for BOLL (exploratory). ') + '</w:t></w:r><w:r>'
      '<w:t xml:space="preserve">'
      + esc('Comparing the five pre-declared candidate rules (Methods; Supplementary Note '
            'S2.13), the region-mean rule was selected in 65.1% of bootstrap replicates and '
            '100% of benign leave-one-out folds (anchor-only 34.9%; OR, AND and partner-only '
            'never selected), retaining 92.9% premenopausal tumour positivity with 2.0% benign '
            'surgical positivity, versus 45.8% under the OR rule (159/347, partner-driven). '
            'These frequencies measure internal rule-selection stability, not independent '
            'validation; the rule and its threshold remain to be evaluated and locked in the '
            'independent development cohort.')
      + '</w:t></w:r></w:p>')
anchor4 = 'in an independent threshold-setting cohort containing symptomatic benign samples before any validation.'
n = xml.count(anchor4)
assert n == 1, 'R4 anchor %d' % n
i = xml.find(anchor4)
e = xml.find('</w:p>', i)
xml = xml[:e + len('</w:p>')] + R4 + xml[e + len('</w:p>'):]
print('R4 inserted')

# ---- R13: Table 1 rows (clone GSE223817 row) ----
def clone_row_with_cells(xml, row_marker, new_rows, tag):
    m = re.search(r'<w:tr\b[^>]*>(?:(?!</w:tr>).)*?' + re.escape(row_marker) + r'.*?</w:tr>', xml, re.S)
    assert m, tag + ': row not found'
    block = m.group(0)
    out = ''
    for cells in new_rows:
        nb = block
        ts = list(re.finditer(r'<w:t\b[^>]*>.*?</w:t>', nb, re.S))
        assert len(ts) >= len(cells), tag + ': cells mismatch %d' % len(ts)
        # replace last-to-first to keep offsets valid
        for (txt, tm) in zip(reversed(cells), reversed(ts[:len(cells)])):
            nb = nb[:tm.start()] + '<w:t xml:space="preserve">' + esc(txt) + '</w:t>' + nb[tm.end():]
        out += nb
    return xml[:m.end()] + out + xml[m.end():]

xml = clone_row_with_cells(xml, 'GSE223817', [
  ['GSE87621', '450K', '5 benign endometrial samples',
   'Background: supportive benign (first direct ZSCAN12 assessment; S2.11)'],
  ['GSE73948', '450K', '26 eutopic endometrium (endometriosis, secretory phase)',
   'Background: supportive benign (S2.11)'],
], 'R13')
print('R13 Table 1 rows added')

with zipfile.ZipFile(DST, 'w', zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == 'word/document.xml':
            data = xml.encode('utf-8')
        zout.writestr(item, data)
zin.close()

z = zipfile.ZipFile(DST)
assert z.testzip() is None
ET.fromstring(z.read('word/document.xml'))
x = z.read('word/document.xml').decode('utf-8')
checks = ['did not cause either lead to fail', 'post hoc exploratory region-level rules',
          'most frequently favoured the region-mean rule', 'most frequently selected over OR',
          'may be optimistically biased', 'we identified no eligible publicly available',
          'does not constitute independent validation of the strict funnel',
          'rather than a settled protocol', 'reproduces the main analysis exactly',
          'rank 2' + MD + '6', 'direction and magnitude of background bias',
          'Region-level rules for BOLL (exploratory).', 'GSE87621', 'GSE73948',
          'Supplementary Note S2.17']
for k in checks:
    print(k[:50], '->', k in x)
print('residual upper bounds:', x.count('upper bounds'))
print('residual influenced either lead:', x.count('influenced either lead'))
print('residual likely overestimates:', x.count('likely overestimates'))
print('residual no public 450K/EPIC cohorts exist:', x.count('no public 450K/EPIC cohorts exist'))
print('residual stably selected:', x.count('stably selected'))
import docx
d = docx.Document(DST)
print('tables:', len(d.tables), '| table1 rows:', len(d.tables[0].rows))
print('saved:', DST)
