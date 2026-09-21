# -*- coding: utf-8 -*-
"""v9.16: 10th review round — (1) GSE93589 methods sentence; (2) BOLL expression
methods paragraph; (3) 'panel' terminology unification (candidate sets /
experimental-pool probes / prioritised candidates; 'panel' only for external
panels, negative claims and the future locked panel); (4) ALL author lines
removed per user instruction. Abstract: menopause wording + upper-bounds scope,
kept <=350 words."""
from docx import Document

SRC = r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.15.docx"
DST = r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.16.docx"

doc = Document(SRC)
log = []

def set_text(p, new):
    for r in list(p.runs):
        r.text = ""
    if p.runs:
        p.runs[0].text = new
    else:
        p.add_run(new)

def replace_in_para(anchor, old, new, must=True):
    n = 0
    for p in doc.paragraphs:
        if anchor in p.text and old in p.text:
            set_text(p, p.text.replace(old, new))
            n += 1
    log.append(("OK " if n >= 1 else "MISS ") + old[:60])
    if must and n == 0:
        raise SystemExit("MISS: " + old[:80])

def replace_para_by_prefix(prefix, new_text):
    n = 0
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            set_text(p, new_text)
            n += 1
    log.append(("OK " if n == 1 else "MISSx%d " % n) + prefix[:50])
    if n != 1:
        raise SystemExit("PREFIX MISS: " + prefix[:60])

def append_to_para(anchor, addition):
    n = 0
    for p in doc.paragraphs:
        if anchor in p.text and addition[:40] not in p.text:
            set_text(p, p.text.rstrip() + " " + addition)
            n += 1
    log.append(("OK " if n >= 1 else "MISS ") + ("APPEND " + anchor[:50]))
    if n == 0:
        raise SystemExit("APPEND MISS: " + anchor[:60])

# ---------- 4. Remove ALL author / affiliation / corresponding lines ----------
removed = 0
for p in list(doc.paragraphs):
    t = p.text.strip()
    if t in ("Dong Liu1, Yubing Chen1*",
             "1 Zhongshan School of Medicine, Sun Yat-sen University, Guangzhou, China",
             "* Corresponding author. Email: liudong1984love@163.com"):
        p._element.getparent().remove(p._element)
        removed += 1
log.append(("OK " if removed == 3 else "MISSx%d " % removed) + "author lines removed")
if removed != 3:
    raise SystemExit("AUTHOR REMOVAL MISS")

# ---------- Abstract rewrite (two recommended changes + terminology + trims) ----------
replace_para_by_prefix("Background: Methylation-based",
"Background: Methylation-based endometrial cancer (EC) detection is maturing rapidly, with an NMPA-approved three-gene kit, international programmes (WID-qEC; CISENDO) and a 2026 Chinese expert consensus. However, existing markers were developed predominantly in postmenopausal women, and marker background is rarely assessed in symptomatic populations or sampling compartments in which tests are deployed. We aimed to build and triage methylation candidates for premenopausal EC.")

replace_para_by_prefix("Methods: We re-analysed",
"Methods: We re-analysed 450K data from TCGA-UCEC and GSE67116 with harmonised differential-methylation criteria and region merging. Candidates were triaged against five background dimensions spanning healthy endometrium, benign endometrium from a surgical cohort used as an intended-use proxy (menopausal status and AUB presentation unavailable), menstrual-cycle phases, cervical scrapes and purified blood cells. Candidate sets were assembled by redundancy pruning and greedy coverage; NMPA-kit genes and WID-qEC components were audited through the same funnel. Sensitivity analyses comprised menopause-stratified, patient-matched, subtype-restricted and adjusted models, cross-preprocessing checks and bootstrap intervals.")

replace_para_by_prefix("Results: Background estimates",
"Results: Background estimates were descriptively higher in the intended-use-proxy cohort than in a cancer-free endometrial reference cohort. cg24589459 (BOLL) was the only probe evaluable across all five background dimensions that met the prespecified gate by point estimate, although its bootstrap confidence interval crossed the threshold (borderline). cg27577527 (ZSCAN12) passed the four probe-level dimensions with the most consistently low background; the intended-use-proxy dimension is not directly assessable (probe absent from the EPIC platform), with only indirect region-level support. Corroboration in independent tumour-tissue cohorts—not the intended sampling compartments—was partial and yielded a post hoc two-probe region-level hypothesis for BOLL. No statistically significant menopause-associated differences were detected, although the analysis was underpowered and did not establish equivalence. Background estimates were pipeline-dependent, so β thresholds are heuristic. Discovery-cohort positivity and coverage estimates are optimistic upper bounds; the candidates are not a validated panel.")

replace_para_by_prefix("Conclusions: Marker background",
"Conclusions: Marker background is cohort- and compartment-dependent; benign controls should come from the target clinical population. BOLL and ZSCAN12 are prioritised for prospective evaluation in premenopausal abnormal-uterine-bleeding triage, where the intended target condition comprises endometrial cancer and atypical hyperplasia; performance for atypical hyperplasia cannot be established from the available datasets. Population matching was approximated, not achieved; confirmation requires a purpose-built premenopausal AUB cohort.")

# ---------- 1. GSE93589 methods sentence ----------
append_to_para("Probe-level corroboration was performed in three independent EPIC cohorts",
    "In addition, GSE93589, a small independent 450K tumour-tissue cohort (nine EC samples; GMQN-normalised β values from the NGDC EWAS Data Hub; not used in discovery), was analysed solely for exploratory corroboration of cg27577527, which is absent from the EPIC manifest.")

# ---------- 2. BOLL expression methods paragraph (inserted after corroboration methods) ----------
idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("Probe-level corroboration was performed"):
        idx = i
        break
assert idx is not None
next_p = doc.paragraphs[idx + 1]
newp = next_p.insert_paragraph_before(
    "BOLL methylation–expression analysis. TCGA-UCEC RNA-seq data (HiSeqV2; RSEM-derived, log2(x+1)-transformed values) were obtained from UCSC Xena and matched to HM450 β values by sample barcode, yielding 172 tumours with paired methylation and expression data. BOLL expression was summarised by the median and by the fraction of samples with expression >1 in tumours and adjacent normals, and methylation–expression association was tested by Spearman's rank correlation.")
log.append("OK expression methods paragraph inserted")

# ---------- 3. 'panel' terminology unification ----------
replace_in_para("auditable funnel",
    "redundancy pruning and greedy panel assembly",
    "redundancy pruning and greedy candidate-set assembly")
replace_in_para("auditable funnel",
    "benchmarking the final candidates against",
    "benchmarking the prioritised candidates against")
replace_para_by_prefix("Panel assembly, sampling-compartment suitability",
    "Candidate-set assembly, sampling-compartment suitability and literature benchmarking")
replace_in_para("greedy incremental coverage",
    "panels were assembled by greedy incremental coverage",
    "candidate sets were assembled by greedy incremental coverage")
replace_in_para("Probes were harmonised by exact probe ID",
    "21 of 23 panel probes are present on the EPIC manifest",
    "21 of 23 experimental-pool probes are present on the EPIC manifest")
replace_in_para("Of the 23 panel probes",
    "Of the 23 panel probes, 21 are present",
    "Of the 23 experimental-pool probes, 21 are present")
replace_in_para("quantitative methylation gradient",
    "whether panel loci show",
    "whether experimental-pool loci show")
replace_in_para("quantitative methylation gradient",
    "for all 21 evaluable panel probes",
    "for all 21 evaluable experimental-pool probes")
replace_in_para("No wet-lab validation has been performed",
    "the proposed panel is designed for",
    "the proposed experimental pool is designed for")
replace_para_by_prefix("Table 5. Literature convergence of panel genes",
    "Table 5. Literature convergence of candidate genes.")
replace_in_para("Menstrual-cycle drift of the anchor probe",
    "full-panel range 0.014–0.042",
    "experimental-pool range 0.014–0.042")
# terminology paragraph: reserve 'panel' explicitly
append_to_para("Terminology used throughout",
    "The term 'panel' is reserved for external published panels and for a future locked diagnostic panel, which does not yet exist.")
# table cell
fixed = 0
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                if p.text.strip() == "All EPIC-evaluable panel probes (n=21)":
                    set_text(p, "All EPIC-evaluable experimental-pool probes (n=21)")
                    fixed += 1
log.append(("OK " if fixed == 1 else "MISSx%d " % fixed) + "table cell panel probes")

doc.save(DST)
print("\n".join(log))
print("saved", DST)

d2 = Document(DST)
abstract = [p.text for p in d2.paragraphs if p.text.strip().startswith(("Background:", "Methods:", "Results:", "Conclusions:"))][:4]
print("ABSTRACT WORDS:", sum(len(t.split()) for t in abstract), [len(t.split()) for t in abstract])
print("TOTAL WORDS:", len("\n".join(p.text for p in d2.paragraphs if p.text.strip()).split()))
