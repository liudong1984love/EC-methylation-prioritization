"""v9.11 verification checklist."""
from docx import Document
doc = Document('Manuscript_EC_methylation_panel_v9.11.docx')
text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
tables_text = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
full = text + '\n' + tables_text
words = len(text.split())
print(f'WORD COUNT: {words} (v9.10 was 6474)')

checks = [
    # repositioning
    ('title informed', 'Population- and compartment-informed in silico prioritization' in text),
    ('no -aware anywhere', 'compartment-aware' not in full),
    ('no "External validation" headings/cells', 'External validation' not in full.replace('external validation section', '')),
    ('corroboration methods heading', '\nIndependent tissue-cohort corroboration\n' in text),
    ('corroboration results heading', 'Independent tissue-cohort corroboration in EPIC cohorts' in text),
    # BOLL/ZSCAN12 asymmetry
    ('no five-dim claim ZSCAN12', 'do not claim five-dimension compliance for ZSCAN12' in text),
    ('new range 0.100-0.141', '0.100–0.141' in text),
    ('old range gone', '0.087–0.141' not in text),
    # post hoc
    ('post hoc rule label', 'post hoc, hypothesis-generating rule derived from these same data' in text),
    ('post hoc abstract', 'post hoc two-probe region-level hypothesis' in text),
    # intended use
    ('intended use definition', 'target condition: EC and atypical hyperplasia; reference standard: histopathology' in text),
    ('intended use intro', 'triage of premenopausal women with abnormal uterine bleeding' in text),
    # heuristic thresholds
    ('heuristic in methods', 'heuristic screening rules for candidate prioritisation' in text),
    ('P90/max reporting', 'maximum, P90 and the proportion of samples above 0.15' in text),
    # new analyses
    ('cross-pipeline para', 'Cross-pipeline consistency.' in text),
    ('median shift number', 'median of 0.022 (maximum 0.045)' in text),
    ('11 of 26', '11 of 26 probes between noob and GMQN' in text),
    ('spearman 0.98', 'Spearman ρ=0.98' in text),
    ('adjusted models', 'age-plus-histology-adjusted limma models' in text),
    ('max dbeta change 0.03', 'maximum change in Δβ 0.03' in text),
    ('subtype F-test', 'subtype F-test FDR<0.05 for all 26 probes' in text),
    # numeric fixes
    ('new bootstrap CI', '0.069–0.182' in text),
    ('old CI gone', '0.099–0.203' not in text),
    ('substitute pipeline caveat', 'P95 0.160–0.167' in text),
    ('no 0.149v0.151 overread', '0.149 versus 0.151) as meaningful' in text),
    # preserved key claims
    ('borderline BOLL kept', 'borderline' in text),
    ('P<=0.002 kept', 'P≤0.002' in text),
    ('upper bounds kept', 'upper bounds' in text),
    ('not validated panel kept', 'not a validated panel' in text),
    ('23-probe pool kept', '23-probe experimental' in text or '23-probe' in text),
    ('table cells renamed', 'Tissue-cohort corroboration' in tables_text),
    ('table8 caption', 'Probe-level corroboration in independent EPIC tumour-tissue cohorts' in text),
]
fails = 0
for name, ok in checks:
    print(('PASS ' if ok else 'FAIL ') + name)
    fails += (not ok)
print(f'\n{fails} FAILURES')

# residual risk spots
import re
for pat in ['0.087', '0.1314', '0.1558', 'five-dimension union-compliant', 'illustrative, non-locked']:
    hits = [l[:130] for l in full.split('\n') if pat in l]
    print(f'-- residual "{pat}": {len(hits)}')
    for h in hits[:3]: print('   ', h)
