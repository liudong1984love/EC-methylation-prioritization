"""v9.4 -> v9.5: implement Li Yunhui review comments #3-#8 (English manuscript).
#3 threshold rationale; #4 Zhou re-annotation probe uniqueness; #5 hyperplasia gradient;
#6 BOLL expression-methylation; #7 cycle drift (Table 4 column + text); #8 concrete future directions.
"""
import csv
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Inches

BASE = Path(r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19")
SRC = BASE / "Manuscript_EC_methylation_panel_v9.4.docx"
DST = BASE / "Manuscript_EC_methylation_panel_v9.5.docx"

drift = {r["probe"]: r["cycle_drift"] for r in csv.DictReader(open(BASE / "results/boll_cycle_drift_v95.csv", encoding="utf-8"))}

doc = Document(SRC)
ps = doc.paragraphs

def append_run(p, text):
    """Append text as a new run, cloning formatting of the last existing run."""
    run = p.add_run(text)
    if p.runs and len(p.runs) > 1:
        prev = p.runs[-2]._r.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
        if prev is not None:
            run._r.insert(0, deepcopy(prev))
    return run

targets = {}
for i, p in enumerate(ps):
    t = p.text.strip()
    if t.startswith("Candidates were triaged against five dimensions"):
        targets["triage"] = p
    elif t.startswith("IDATs were processed with minfi"):
        targets["preproc"] = p
    elif t.startswith("Two probes passed all measurable background dimensions"):
        targets["final"] = p
    elif t.startswith("BOLL (cg24589459) is, to our knowledge"):
        targets["boll_disc"] = p
    elif t.startswith("Several limitations warrant emphasis"):
        targets["limitations"] = p
    elif t.startswith("In conclusion, multi-cohort discovery"):
        targets["conclusion"] = p
    elif t == "Discussion":
        targets["discussion_h"] = p
    elif t == "External validation":
        targets["ev_heading"] = p
    elif t.startswith("Probe-level external validation was performed"):
        targets["ev_body"] = p
    elif t.startswith("Results: TCGA and GSE67116 yielded"):
        targets["abstract_res"] = p
    elif t.startswith("*Not assessable in EPIC-based cohorts"):
        targets["tab4_note"] = p
    elif t.startswith("Table 5. Literature convergence"):
        targets["tab5_caption"] = p
    elif t.startswith("NMPA."):
        targets["last_ref"] = p

missing = [k for k in ("triage", "preproc", "final", "boll_disc", "limitations", "conclusion",
                       "discussion_h", "ev_heading", "ev_body", "abstract_res", "tab4_note", "tab5_caption", "last_ref") if k not in targets]
if missing:
    raise SystemExit(f"missing anchors: {missing}")

# --- #3 threshold rationale (Methods, triage para) ---
append_run(targets["triage"],
    " The 0.15 gate was chosen as an intermediate value accommodating both healthy-tissue and "
    "symptomatic or cervical sampling backgrounds; candidate rankings were robust across gates of "
    "0.10\u20130.20 (Results, threshold sensitivity), and the \u0394\u03b2\u22650.30 entry criterion matches "
    "effect sizes reported for validated gynaecologic methylation markers.")

# --- #4 probe uniqueness (Methods, preprocessing para) ---
append_run(targets["preproc"],
    " Probe identities were additionally cross-checked against the Zhou et al. hg38 re-annotation of "
    "the HM450 array [28]: the two BOLL-region probes (cg24589459, type II; cg07495363, type I) each map "
    "uniquely to the BOLL locus (chr2:197,786,574 and chr2:197,786,304, respectively; mapping quality 60 "
    "for both), excluding cross-hybridization as an explanation for their discordant behaviour.")

# --- #7 cycle drift sentence (Results, final candidates para) ---
append_run(targets["final"],
    " Menstrual-cycle drift of the anchor probe was minimal (mean |\u0394\u03b2| 0.026 between "
    "fingerprint-paired early- and mid-secretory samples; 0.042 for cg07495363; full-panel range "
    "0.014\u20130.042; Table 4), excluding cycle synchronisation as an explanation for the low benign "
    "background of BOLL.")

# --- abstract: gradient highlight ---
append_run(targets["abstract_res"],
    " A quantitative hyperplasia-to-carcinoma methylation gradient was observed at panel loci in two cohorts.")

# --- #5 + #6 new Results subsection, inserted before the Discussion heading ---
h = targets["discussion_h"].insert_paragraph_before(
    "Hyperplasia-to-carcinoma gradient and BOLL expression analysis")
h.style = targets["ev_heading"].style

b1 = targets["discussion_h"].insert_paragraph_before(
    "We next tested whether panel loci show a quantitative methylation gradient across the neoplastic "
    "continuum. In GSE136791 (EPIC), mean \u03b2 increased monotonically from benign endometrium (the 13 "
    "endometrial controls of GSE155760, same platform, used as a cross-cohort benign reference) through "
    "hyperplasia (n=34) to carcinoma (n=69) for all 21 evaluable panel probes (median mean \u03b2 0.152, "
    "0.601 and 0.665, respectively). For the BOLL anchor cg24589459, mean \u03b2 rose from 0.042 in controls "
    "to 0.392 in hyperplasia and 0.472 in carcinoma (positivity at \u03b2>0.3: 67.6% versus 82.6%); its "
    "partner cg07495363 showed the same ordering (0.206, 0.638, 0.680). The within-cohort "
    "hyperplasia<carcinoma ordering held for 21/21 probes in GSE136791 and was reproduced in GSE67116 "
    "(450K; 8 hyperplasia versus 33 primary EC), where all 23 panel probes showed higher mean \u03b2 in "
    "carcinoma (median 0.533 versus 0.327; BOLL cg24589459 0.454 versus 0.295). Because GSE67116 "
    "contributed to discovery it does not constitute independent evidence, and because hyperplasia "
    "histology (atypical versus non-atypical) is not annotated in either cohort, these data support but "
    "do not prove detection of a premalignant continuum. They nevertheless motivate quantitative, "
    "region-level assay readouts that could stratify risk rather than merely flag carcinoma.")
b1.style = targets["ev_body"].style

b2 = targets["discussion_h"].insert_paragraph_before(
    "To probe the biology of the BOLL signal, we correlated methylation with expression in 172 TCGA-UCEC "
    "tumours with paired HM450 methylation and RNA-seq data. Methylation at cg24589459 correlated weakly "
    "and inversely with BOLL expression (Spearman \u03c1=\u22120.19, p=0.011; cg07495363 \u03c1=\u22120.17, "
    "p=0.023). However, BOLL expression was essentially absent from endometrial tissue irrespective of "
    "methylation state (median 0 in both tumours and adjacent normals; above-background expression in "
    "15.9% of tumours and 0% of adjacent normals), whereas mean \u03b2 was 0.637 in tumours versus 0.098 "
    "in adjacent normals. BOLL hypermethylation in EC therefore does not silence an actively expressed "
    "gene and is best interpreted as ectopic hypermethylation of a germline-specific locus that is "
    "already repressed in benign endometrium.")
b2.style = targets["ev_body"].style

# --- #4 interpretation + #6 biology bounds (Discussion, BOLL para) ---
append_run(targets["boll_disc"],
    " Two additional observations clarify the nature of the BOLL signal. First, both BOLL-region probes "
    "map uniquely to the locus [28], so the dissociation between the anchor probe (low background, "
    "moderate positivity) and its partner probe (higher background, high sensitivity), 270 bp away, most "
    "likely reflects genuine regional methylation heterogeneity rather than probe artefact\u2014reinforcing "
    "the need for region-level assay design. Second, the expression analysis above indicates that BOLL is "
    "a passenger rather than a driver marker: we found no evidence that its hypermethylation silences an "
    "expressed tumour suppressor, and we make no claim of pathway involvement (for example PI3K/AKT or "
    "oestrogen signalling). Passenger status does not reduce diagnostic utility\u2014many clinically used "
    "methylation markers are passengers\u2014but it bounds the biological interpretation.")

# --- #8 concrete future directions, inserted before the conclusion ---
fut = targets["conclusion"].insert_paragraph_before(
    "These limitations define a concrete validation agenda. Wet-lab validation is planned as a two-stage "
    "targeted bisulfite study in premenopausal women with abnormal uterine bleeding, using multiplex qMSP "
    "for binary calls and droplet-digital MSP or targeted bisulfite sequencing for quantitative, "
    "region-level readouts. Critically, the benign comparator arm must span the realistic differential "
    "diagnosis of premenopausal bleeding\u2014endometrial polyps, leiomyomas, adenomyosis and anovulatory "
    "cycles\u2014rather than healthy donors, because symptomatic-benign background is the decisive "
    "determinant of specificity. A pilot of approximately 50 carcinomas and 100 benign controls would "
    "estimate sensitivity (expected \u224890%) and specificity (expected \u226585%) with 95% confidence "
    "half-widths of \u22488% and \u22487%, respectively; a definitive cohort of approximately 140 "
    "carcinomas and 200 benign controls would narrow both to \u00b15%. Even 20 samples per group would "
    "provide >99% power to distinguish 90% tumour positivity from a 30% benign background-positive "
    "fraction under the illustrative \u03b2>0.3 rule.")
fut.style = targets["limitations"].style

# --- #7 Table 4: add Cycle drift column ---
tab4 = None
for t in doc.tables:
    if "Tier" in " ".join(c.text for c in t.rows[0].cells):
        tab4 = t
        break
if tab4 is None:
    raise SystemExit("Table 4 not found")
hdr = [c.text.strip() for c in tab4.rows[0].cells]
probe_col = next(i for i, htxt in enumerate(hdr) if htxt.lower().startswith("probe"))
tab4.add_column(Inches(0.9))
tab4.rows[0].cells[-1].text = "Cycle drift |\u0394\u03b2|"
n_filled = 0
for row in tab4.rows[1:]:
    probe = row.cells[probe_col].text.strip().split()[0].rstrip("*")
    val = drift.get(probe, "\u2014")
    row.cells[-1].text = val
    n_filled += 1
print(f"Table 4: drift column filled for {n_filled} rows")

# --- Table 4 footnote for the new column ---
note = targets["tab5_caption"].insert_paragraph_before(
    "Cycle drift: mean |\u0394\u03b2| between fingerprint-paired early- and mid-secretory endometrium "
    "(GSE90060); smaller values indicate higher robustness to menstrual-cycle phase.")
note.style = targets["tab4_note"].style

# --- reference [28] ---
all_ps = doc.paragraphs
idx_last = next(i for i, p in enumerate(all_ps) if p._p is targets["last_ref"]._p)
ref_text = ("Zhou W, Laird PW, Shen H. Comprehensive characterization, annotation and innovative use of "
            "Infinium DNA methylation BeadChip probes. Nucleic Acids Res. 2017;45(4):e22.")
if idx_last + 1 < len(all_ps):
    new_ref = all_ps[idx_last + 1].insert_paragraph_before(ref_text)
else:
    new_ref = doc.add_paragraph(ref_text)
new_ref.style = targets["last_ref"].style

doc.save(DST)
print("saved:", DST)
