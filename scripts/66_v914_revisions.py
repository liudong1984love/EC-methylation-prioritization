# -*- coding: utf-8 -*-
"""v9.14: 8th review round — abstract compression, citation mapping, target-condition
consistency, dbeta sourcing, adjusted-model details, pipeline-scope limit, title page."""
from docx import Document

SRC = r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.13.docx"
DST = r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19\Manuscript_EC_methylation_panel_v9.14.docx"

doc = Document(SRC)
log = []

def set_text(p, new):
    for r in list(p.runs):
        r.text = ""
    if p.runs:
        p.runs[0].text = new
    else:
        p.add_run(new)

def replace_in_para(anchor, old, new, must=True):
    n = 0
    for p in doc.paragraphs:
        if anchor in p.text and old in p.text:
            set_text(p, p.text.replace(old, new))
            n += 1
    log.append(("OK " if n >= 1 else "MISS ") + old[:60])
    if must and n == 0:
        raise SystemExit("MISS: " + old[:80])

def replace_para_by_prefix(prefix, new_text):
    n = 0
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            set_text(p, new_text)
            n += 1
    log.append(("OK " if n == 1 else "MISSx%d " % n) + prefix[:50])
    if n != 1:
        raise SystemExit("PREFIX MISS: " + prefix[:60])

def append_to_para(anchor, addition):
    n = 0
    for p in doc.paragraphs:
        if anchor in p.text and addition[:40] not in p.text:
            set_text(p, p.text.rstrip() + " " + addition)
            n += 1
    log.append(("OK " if n >= 1 else "MISS ") + ("APPEND " + anchor[:50]))
    if n == 0:
        raise SystemExit("APPEND MISS: " + anchor[:60])

# ---------- 0. Title page: authors / affiliation / corresponding ----------
for p in doc.paragraphs:
    if p.text.strip() == "Abstract":
        p.insert_paragraph_before("Dong Liu1, Yubing Chen1*")
        p.insert_paragraph_before("1 Zhongshan School of Medicine, Sun Yat-sen University, Guangzhou, China")
        p.insert_paragraph_before("* Corresponding author. Email: liudong1984love@163.com")
        log.append("OK title page authors inserted")
        break

# ---------- 1. Abstract rewrite (575 -> ~355 words) ----------
replace_para_by_prefix("Background: Methylation-based",
"Background: Methylation-based endometrial cancer (EC) detection is maturing rapidly, with an NMPA-approved three-gene adjunctive kit, international programmes (WID-qEC; CISENDO) and a 2026 Chinese expert consensus. However, existing markers were discovered and validated predominantly in postmenopausal women, and marker background is rarely assessed in the symptomatic populations or sampling compartments in which tests are deployed. We aimed to build and rigorously triage methylation candidates for validation in premenopausal EC populations.")

replace_para_by_prefix("Methods: We re-analysed",
"Methods: We re-analysed 450K data from TCGA-UCEC and GSE67116 with harmonised differential-methylation criteria and region merging. Candidates were triaged against five background dimensions spanning healthy and benign endometrium (the latter approximating the intended-use premenopausal population), menstrual-cycle phases, cervical scrapes and purified blood cells. Panels were assembled by redundancy pruning and greedy coverage; the NMPA-kit genes and WID-qEC components were audited through the same funnel. Sensitivity analyses comprised menopause-stratified, patient-matched and subtype-restricted analyses, confounder-adjusted models, cross-preprocessing checks and bootstrap confidence intervals.")

replace_para_by_prefix("Results: TCGA and GSE67116 yielded",
"Results: Background estimates were higher in the intended-use-proxy cohort than in a cancer-free endometrial reference cohort. Among the lead candidates, cg24589459 (BOLL) was the only probe that both could be evaluated directly across all five background dimensions and met the prespecified gate by point estimate, although its bootstrap confidence interval crossed the threshold (borderline). cg27577527 (ZSCAN12) passed the four probe-level dimensions with the most consistently low background; the intended-use-proxy dimension is not directly assessable because the probe is absent from the EPIC platform, and support there is indirect and region-level only. Corroboration in independent tumour-tissue cohorts—not the intended sampling compartments—partially supported the candidates, with a post hoc region-level hypothesis for BOLL. No significant menopause-associated differences were detected, although the analysis was underpowered. Background estimates were pipeline-dependent, so the fixed β thresholds are heuristic screening rules. Estimates derive partly from discovery data and are upper bounds; the candidates are not a validated panel.")

replace_para_by_prefix("Conclusions: Marker background",
"Conclusions: Marker background is cohort- and compartment-dependent; benign controls should come from the target clinical population. BOLL and ZSCAN12 are prioritised as candidates for detecting endometrial cancer and atypical hyperplasia in premenopausal abnormal-uterine-bleeding triage; they do not constitute a validated panel, and population matching was approximated rather than achieved. Confirmation requires wet-lab validation in a purpose-built premenopausal AUB cohort.")

# ---------- 2. Citation mapping fixes ----------
replace_in_para("obtained NMPA approval",
    "obtained NMPA approval as an adjunctive diagnostic for suspected EC, and a 2026 expert consensus provided guidance on the clinical application of methylation testing in China [9,10,27]",
    "obtained NMPA approval as an adjunctive diagnostic for suspected EC [27], and a 2026 expert consensus provided guidance on the clinical application of methylation testing in China [10]")
replace_in_para("CDO1 anchors the NMPA-approved kit",
    "CDO1 anchors the NMPA-approved kit [10]",
    "CDO1 anchors the NMPA-approved kit [27]")
replace_in_para("the anchor of the only approved national kit",
    "(CDO1 [8,10])", "(CDO1 [8,27])")
replace_in_para("entered clinical reality",
    "with an approved kit in China, a consensus statement, and international programmes approaching routine triage [5,6,8,9,10]",
    "with an approved kit in China [27], a consensus statement [10], and international programmes approaching routine triage [5,6,8,9]")

# ---------- 3. Target-condition consistency (EC + AH) ----------
replace_para_by_prefix("In conclusion, multi-cohort discovery",
"In conclusion, multi-cohort discovery with population- and compartment-informed background triage prioritised the BOLL and ZSCAN12 regions for further evaluation in the detection of endometrial cancer and atypical hyperplasia during premenopausal AUB triage, and reconciled the composition of existing tests with their sampling designs. These findings do not establish diagnostic performance or suitability for self-collected specimens; a locked region-level assay and thresholds should be prospectively validated in a purpose-built premenopausal AUB cohort containing representative benign uterine conditions. We distinguish three outputs of this study: two lead candidate regions (BOLL and ZSCAN12), a 23-probe experimental pool for wet-lab adjudication, and a final locked diagnostic panel, which does not yet exist.")

append_to_para("avoid carcinoma-specificity claims",
    "Separate performance estimates for atypical versus non-atypical hyperplasia therefore cannot be computed from these public data.")

append_to_para("comb-p-like algorithm",
    "In GSE67116 the only available disease contrast was carcinoma versus hyperplasia (atypia unannotated); DMRs from this cohort therefore represent carcinoma-versus-hyperplasia differences and cannot be attributed to atypical versus non-atypical lesions.")

# ---------- 4. dbeta source annotation ----------
replace_in_para("Δβ=0.533", "borderline), Δβ=0.533,",
    "borderline), Δβ=0.533 (TCGA tumour versus adjacent normal),")
replace_in_para("Δβ=0.614", "(P95 β 0.044–0.060; Δβ=0.614; positivity 92.9%)",
    "(P95 β 0.044–0.060; Δβ=0.614, same contrast; positivity 92.9%)")
replace_in_para("elite single probes were defined by",
    "Δβ≥0.50, tumour positivity", "Δβ≥0.50 (TCGA tumour versus adjacent normal), tumour positivity")

# ---------- 5. Adjusted-model Methods details ----------
append_to_para("P values were Benjamini–Hochberg adjusted",
    "Covariate-adjusted sensitivity models used limma on M-values with complete-case design matrices: group plus age (n=462 of 477; 15 samples lacking age excluded), group plus age and histological type (endometrioid versus other; n=462), an endometrioid-only age-adjusted model (309 endometrioid tumours and 34 adjacent normals with complete covariates), and a tumour-only molecular-subtype model (n=398 of 431 tumours with subtype annotation).")

# ---------- 6. Cross-pipeline scope limit ----------
replace_in_para("Candidate ranking is thus stable",
    "Candidate ranking is thus stable across mainstream pipelines, but absolute pass/fail status near the gate is pipeline-bound",
    "Within the two datasets and three preprocessing routes tested here, candidate ranking was stable, but absolute pass/fail status near the gate was pipeline-bound")

# ---------- 7. Compression: limitations + cohort-dependence paragraph ----------
replace_in_para("Discovery and triage layers used different normalisation",
    "Discovery and triage layers used different normalisation (noob versus GMQN); near-gate pass/fail status changed for 11 of 26 probes between pipelines (both lead probes robust), so thresholds are heuristic and pipeline-bound.",
    "Discovery and triage layers used different normalisation (noob versus GMQN), and near-gate pass/fail status was pipeline-bound (Results), so the fixed thresholds are heuristic.")
replace_in_para("k=0 of B=500 permutations",
    "(chromosome-wise empirical P≤0.002, k=0 of B=500 permutations with the +1 correction (k+1)/(B+1)=1/501≈0.002; not a genome-wide FWER)",
    "(chromosome-wise empirical P≤0.002 with the (k+1)/(B+1) correction; not a genome-wide FWER)")

replace_para_by_prefix("The cohort dependence we quantified",
"The cohort dependence we quantified (background inflation in the intended-use-proxy cohort versus healthy endometrium) offers a concrete explanation for the attrition of methylation markers during translation and argues that benign controls must be drawn from the intended-use population. Public data remain devoid of premenopausal benign AUB methylation profiles (polyps, leiomyomas, adenomyosis), a gap only purpose-built cohorts can fill.")

doc.save(DST)
print("\n".join(log))
print("saved", DST)

# abstract word count
d2 = Document(DST)
abstract = [p.text for p in d2.paragraphs if p.text.strip().startswith(("Background:", "Methods:", "Results:", "Conclusions:"))][:4]
total = sum(len(t.split()) for t in abstract)
print("ABSTRACT WORDS:", total, [len(t.split()) for t in abstract])
allw = len("\n".join(p.text for p in d2.paragraphs if p.text.strip()).split())
print("TOTAL WORDS:", allw)
