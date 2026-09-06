"""Apply review-consensus fixes -> v9. Includes recomputation from local data."""
import csv
import re
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document

BASE = Path(r"C:\Users\ld\WorkBuddy\2026-08-19-15-50-19")
DATA = BASE / "data"
DL = Path(r"C:/Users/ld/Downloads")
SRC = BASE / "Manuscript_EC_methylation_panel_v8.docx"
DST = BASE / "Manuscript_EC_methylation_panel_v9.docx"

# ---------- 数据重算 ----------
panel = [l.split(",")[0] for l in open(BASE / "results/wetlab_panel_23_primer_annotation.csv", encoding="utf-8").readlines()[1:] if l.strip()]
EPIC_ABSENT = {"cg27577527", "cg20998319"}
EVAL = [p for p in panel if p not in EPIC_ABSENT]
need = set(EVAL)

def read_list(fn):
    return [m.group(0) for line in open(DL / fn, encoding="utf-8") for m in [re.search(r"GSM\d+", line)] if m]

ca69 = read_list("NoteGPT_GSE136791_carcinoma_GSM_list.txt")
eec = read_list("NoteGPT_GSE155760_EEC_GSM_list.txt")
ctrl13 = read_list("NoteGPT_GSE155760_endometrial_controls_GSM_list.txt")

def vals(gse, ids):
    out = {p: [] for p in EVAL}
    for g in ids:
        f = DATA / gse / "ewas_betas" / f"{g}.txt"
        if not f.exists():
            continue
        with open(f) as fh:
            for line in fh:
                pr, b = line.rstrip("\n").split("\t")
                if pr in need and b != "NA":
                    out[pr].append(float(b))
    return out

all103 = [f.stem for f in (DATA / "GSE136791/ewas_betas").glob("GSM*.txt")]
v_all, v_ca, v_eec, v_ct = vals("GSE136791", all103), vals("GSE136791", ca69), vals("GSE155760", eec), vals("GSE155760", ctrl13)

def count80(v):
    return sum(1 for p in EVAL if v[p] and np.mean(np.array(v[p]) > 0.3) >= 0.8)
def rng80(v):
    r = [np.mean(np.array(v[p]) > 0.3) for p in EVAL if v[p]]
    return min(r), max(r)
def p95range(v):
    r = [float(pd.Series(v[p]).quantile(0.95)) for p in EVAL if v[p]]
    return min(r), max(r)

n_all, n_ca = count80(v_all), count80(v_ca)
rng_ca = rng80(v_ca)
n_eec = count80(v_eec); rng_eec = rng80(v_eec)
ct_rng = p95range(v_ct)
print(f"all103: {n_all}/21>=80%; ca69: {n_ca}/21 ({rng_ca[0]*100:.0f}-{rng_ca[1]*100:.0f}%); eec: {n_eec}/21 ({rng_eec[0]*100:.0f}-{rng_eec[1]*100:.0f}%); ctrl13 P95 {ct_rng[0]:.2f}-{ct_rng[1]:.2f}")

# BOLL OR-rule（示意性）
b1, b2 = "cg24589459", "cg07495363"
def or_rule(v):
    tp = sum(1 for i in range(len(v[b1])) if v[b1][i] > 0.3 or v[b2][i] > 0.3)
    n = len(v[b1])
    return tp, n
tp_ca, n_ca_ = or_rule(v_ca); tp_ct, n_ct_ = or_rule(v_ct)
or_sens = tp_ca / n_ca_ * 100
or_fp = tp_ct / n_ct_ * 100
print(f"BOLL OR-rule: ca69 {tp_ca}/{n_ca_} = {or_sens:.1f}%; ctrl13 FP {tp_ct}/{n_ct_} = {or_fp:.1f}%")

# ---------- 文档修改 ----------
doc = Document(SRC)
ps = doc.paragraphs

def set_text(p, text):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)

def sub_in_para(p, old, new):
    if old in p.text:
        set_text(p, p.text.replace(old, new))
        return True
    return False

# 1) Methods P27
p27 = ps[27]
sub_in_para(p27, "External validation was performed at the probe level in two independent EPIC-array cohorts not used in discovery: GSE136791 (103 endometrial samples; 69 endometrioid carcinoma and 34 hyperplasia, histology annotations per GEO) and GSE155760,",
            "External validation was performed at the probe level in three independent EPIC-array cohorts not used in discovery: GSE136791 (103 endometrial samples; 69 endometrioid carcinoma and 34 hyperplasia, histology annotations per GEO), GSE155760,")
sub_in_para(p27, "Data were obtained from GEO (GSE136791; GSE155760) and processed with GMQN-normalised β values where provided, otherwise noob normalisation of IDATs.",
            "Data were obtained from GEO (GSE136791; GSE155760) and processed with GMQN-normalised β values where provided, otherwise noob normalisation of IDATs. GSE178610 (80 endometrioid EC; 49 fresh-frozen and 31 FFPE) was obtained as submitter-processed β values from the GEO series matrix—processing not homologous to the GMQN/noob pipelines—and was analysed stratified by preservation type.")
sub_in_para(p27, "control background (P95 β) on the cancer-free gynaecologic controls (n=36);",
            "control background (P95 β) primarily on the 13 endometrial-mucosa controls;")

# 2) Results P36 首句
p36 = ps[36]
sub_in_para(p36, "The two fully compliant candidates were cg24589459 (BOLL;",
            "Two probes passed all measurable background dimensions: cg24589459 (BOLL, assessable in all five cohorts;")
sub_in_para(p36, "cg27577527 (ZSCAN12; P95 β 0.044–0.060 where measurable, Δβ=0.614, positivity 92.9%; the probe is absent from EPIC data).",
            "and cg27577527 (ZSCAN12; P95 β 0.044–0.060 in the four measurable dimensions, Δβ=0.614, positivity 92.9%; the probe is absent from EPIC data and was therefore not assessable in the symptomatic-benign EPIC cohort).")

# 3) Table 2 漏斗末行
t2 = doc.tables[1]
last = t2.rows[-1]
set_text(last.cells[0].paragraphs[0], "Passing all measurable background dimensions")
set_text(last.cells[1].paragraphs[0], "2")
r1 = t2.add_row()
set_text(r1.cells[0].paragraphs[0], "— of which assessable in all five background cohorts")
set_text(r1.cells[1].paragraphs[0], "1 (BOLL)")
r2 = t2.add_row()
set_text(r2.cells[0].paragraphs[0], "— symptomatic-benign EPIC cohort not assessable")
set_text(r2.cells[1].paragraphs[0], "1 (ZSCAN12)")

# 4) 样式修正
ps = doc.paragraphs
for p in ps:
    if p.text.startswith("Threshold sensitivity and cross-platform coverage"):
        p.style = doc.styles["Heading 3"]
    if p.text.startswith("All five markers of the internally derived set"):
        p.style = doc.styles["Normal"]

# 5) P45 外部验证段
p45 = None
for p in doc.paragraphs:
    if p.text.startswith("Probe-level external validation was performed in three"):
        p45 = p
        break
sub_in_para(p45, "In GSE136791, 16 of 21 evaluable probes reached at least 80% tumour positivity (β>0.3), including",
            f"In the full 103-sample GSE136791 series, {n_all} of 21 evaluable probes reached at least 80% tumour positivity (β>0.3); restricting to the 69 carcinoma samples, {n_ca} of 21 reached at least 80%, including")
sub_in_para(p45, "a pattern consistent with formaldehyde-fixation-related inflation of methylation signals in FFPE tissue;",
            "a pattern consistent with, but not specific for, fixation- or batch-related inflation of methylation signals in FFPE tissue;")
sub_in_para(p45, "this is consistent with the detection of atypical hyperplasia reported for other methylation markers [25], and implies that tissue-based discrimination between carcinoma and atypical hyperplasia requires quantitative (delta-Ct) rather than threshold-based readouts.",
            "the histological subtype of these hyperplasia samples (atypical versus non-atypical) is not specified in the GEO records, so this pattern is consistent with, but does not prove, detection of premalignant lesions as reported for other methylation markers [25]; it suggests that quantitative rather than binary threshold-based assays will be needed if tissue-based discrimination between carcinoma and hyperplasia is required.")
# 追加示意性 panel 规则句
add_or = (f" As an illustrative, non-locked panel rule, positivity of either BOLL-region probe (β>0.3) yielded "
          f"{or_sens:.1f}% tumour positivity in the 69 GSE136791 carcinomas and a false-positive rate of {or_fp:.1f}% "
          f"in the 13 endometrial controls; formal panel-rule definition and threshold locking are deferred to wet-lab validation.")
set_text(p45, p45.text.rstrip() + add_or)

# 6) P51 ZNF154 引用 + 自我审计软化
for p in doc.paragraphs:
    if "ZNF154, a robust pan-cancer locus [13]" in p.text:
        sub_in_para(p, "ZNF154, a robust pan-cancer locus [13]", "ZNF154, a robust pan-cancer locus [11]")
        set_text(p, p.text.rstrip() + " The set shares the TCGA discovery data with the present analysis and therefore does not constitute independent validation; the audit demonstrates the attrition of tumour-signal-driven candidates under clinically relevant background screening.")
        break

# 7) Table 8 第 9 行：21 探针口径统一
t8 = doc.tables[7]
row9 = t8.rows[9]
set_text(row9.cells[0].paragraphs[0], "All EPIC-evaluable panel probes (n=21)")
set_text(row9.cells[2].paragraphs[0], f"{n_all}/21 ≥80%")
set_text(row9.cells[3].paragraphs[0], f"{n_ca}/21 ≥80% ({rng_ca[0]*100:.0f}–{rng_ca[1]*100:.0f}%)")
set_text(row9.cells[4].paragraphs[0], f"{n_eec}/21 ≥80% ({rng_eec[0]*100:.0f}–{rng_eec[1]*100:.0f}%)")
set_text(row9.cells[5].paragraphs[0], f"P95 {ct_rng[0]:.2f}–{ct_rng[1]:.2f}")
# 增生列（21 探针口径，本地算）
hyp = read_list("NoteGPT_GSE136791_hyperplasia_GSM_list.txt")
v_hyp = vals("GSE136791", hyp)
n_hyp = count80(v_hyp); rng_hyp = rng80(v_hyp)
set_text(row9.cells[7].paragraphs[0], f"{n_hyp}/21 ≥80% ({rng_hyp[0]*100:.0f}–{rng_hyp[1]*100:.0f}%)")

doc.save(DST)
print("saved:", DST)

# 校验
d2 = Document(DST)
t = "\n".join(p.text for p in d2.paragraphs)
for kw in ["[11]", "three independent EPIC-array cohorts", "not specific for", "not specified in the GEO records",
           "illustrative, non-locked panel rule", "shares the TCGA discovery data",
           "In the full 103-sample GSE136791 series"]:
    print(("FOUND  " if kw in t else "MISS   ") + kw)
t8b = d2.tables[7]
print("row9:", [c.text[:26] for c in t8b.rows[9].cells])
